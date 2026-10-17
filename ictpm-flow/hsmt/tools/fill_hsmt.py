#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Ghép 01_hsmt.json (approved) vào Mẫu 5A → 01_hsmt_draft.docx.

Nguyên tắc fill:
- Chỉ ghi đè ô khi JSON có giá trị rõ (Decision/Fact). Thiếu → giữ template, không đoán.
- E-CDNT match đúng mã (không startswith nhầm 3 → 31.4).
- Ô có khung pháp lý (18.2, 5.1, 28.2…): chỉ thay chỗ trống / dòng chọn; giữ phần còn lại.
- Bảng mới: style Table Grid (có kẻ). Văn bản nhiều dòng: mỗi dòng một paragraph.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

MODULE = Path(__file__).resolve().parents[1]
PROJECT = Path(__file__).resolve().parents[3]
JSON_PATH = MODULE / "artifacts" / "01_hsmt.json"
OUT_PATH = MODULE / "artifacts" / "01_hsmt_draft.docx"


def find_template(data: dict) -> Path:
    rel = (data.get("inputs") or {}).get("template")
    if rel:
        cand = Path(rel)
        if not cand.is_absolute():
            cand = PROJECT / rel
        if cand.exists():
            return cand
    folder = MODULE / "templates"
    if folder.exists():
        for p in sorted(folder.glob("*.docx")):
            if "5A" in p.name or "5a" in p.name:
                return p
    raise FileNotFoundError("Không tìm thấy Mẫu 5A tại ictpm-flow/hsmt/templates/")


def _rfonts(run) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    """Ghi text; \\n → nhiều paragraph (xuống hàng thật trong Word)."""
    text = "" if text is None else str(text)
    lines = text.split("\n") if text else [""]

    # Xóa paragraph thừa, giữ lại đủ số dòng
    while len(cell.paragraphs) > len(lines):
        p = cell.paragraphs[-1]._p
        p.getparent().remove(p)
    while len(cell.paragraphs) < len(lines):
        cell.add_paragraph()

    for i, line in enumerate(lines):
        p = cell.paragraphs[i]
        for r in p.runs:
            r.text = ""
        if p.runs:
            p.runs[0].text = line
            p.runs[0].bold = bold
            _rfonts(p.runs[0])
        else:
            run = p.add_run(line)
            run.bold = bold
            _rfonts(run)


def replace_blank(template_text: str, blank_pattern: str, value: str) -> str:
    """Thay chỗ trống kiểu ___ / _____ trong khung template; giữ phần còn lại."""
    if not value:
        return template_text
    return re.sub(blank_pattern, value, template_text, count=1)


def make_paragraph_after(anchor: Paragraph, text: str, bold: bool = False) -> Paragraph:
    """Chèn 1 hoặc nhiều paragraph (tách theo \\n) ngay sau anchor; trả về paragraph cuối."""
    lines = (text or "").split("\n") or [""]
    cur = anchor
    for i, line in enumerate(lines):
    new_p = OxmlElement("w:p")
        cur._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
        run = para.add_run(line)
        run.bold = bold if i == 0 else False
    _rfonts(run)
        cur = para
    return cur


def delete_between(start: Paragraph, end: Paragraph) -> None:
    el = start._p.getnext()
    while el is not None and el is not end._p:
        nxt = el.getnext()
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
        el = nxt


def field_val(obj, *keys, default: str = "") -> str:
    cur = obj
    for k in keys:
        if not isinstance(cur, dict):
            return default
        cur = cur.get(k)
    return as_plain_text(cur, default=default)


def as_plain_text(cur, default: str = "") -> str:
    """Lấy chuỗi điền Word. Nếu là object {value/proposed/...} thì chỉ lấy value (fallback proposed).
    Không bao giờ str(dict) ra Word.
    """
    if cur is None:
        return default
    if isinstance(cur, dict):
        # Tránh dump cả meta confidence/evidence vào ô
        if "value" in cur or "proposed" in cur:
            v = cur.get("value")
            if v is None or (isinstance(v, str) and not v.strip()):
                v = cur.get("proposed")
            if v is None:
                return default
            if isinstance(v, dict):
                return as_plain_text(v, default=default)
            return str(v).strip()
        return default
    return str(cur).strip()


def has_value(fields: dict, key: str) -> bool:
    """True nếu field tồn tại và value không rỗng / không TBD."""
    if key not in fields:
        return False
    v = fields[key]
    if isinstance(v, dict):
        s = str(v.get("value") or "").strip()
    else:
        s = str(v or "").strip()
    if not s:
        return False
    if s.upper().startswith("TBD") or s.lower() in ("null", "none", "?"):
        return False
    return True


def fv(fields: dict, key: str, default: str = "") -> str:
    return field_val(fields, key, default=default)


def find_para(doc: Document, prefix: str, after: Paragraph | None = None) -> Paragraph | None:
    found = None
    started = after is None
    for p in doc.paragraphs:
        if not started:
            if p._p is after._p:
                started = True
            continue
        if p.text.strip().startswith(prefix):
            found = p
    return found


def add_table_row(table: Table):
    tbl = table._tbl
    last = tbl.tr_lst[-1]
    new_tr = deepcopy(last)
    tbl.append(new_tr)
    return table.rows[-1]


def ensure_rows(table: Table, n_data_rows: int, header_rows: int = 2) -> None:
    need = header_rows + n_data_rows
    while len(table.rows) < need:
        add_table_row(table)


def set_table_grid(table: Table) -> None:
    """Gán Table Grid để có kẻ ô."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    # xóa style cũ nếu có
    for child in list(tblPr):
        if child.tag == qn("w:tblStyle"):
            tblPr.remove(child)
    style = OxmlElement("w:tblStyle")
    style.set(qn("w:val"), "TableGrid")
    tblPr.insert(0, style)


def ecdnt_code(cell0: str) -> str:
    """Trả về mã chuẩn: '1.1', '5.1', '7.5', '18.2', '29.3', '3'…"""
    t = re.sub(r"\s+", " ", (cell0 or "").strip())
    m = re.match(r"E-CDNT\s+([\d.]+)", t, flags=re.I)
    if not m:
        return ""
    return m.group(1)


def detect_phuong_phap_tai_chinh(data: dict, fields: dict) -> str:
    """Trả về 'gia_thap_nhat' | 'gia_danh_gia' | ''."""
    chunks = [
        fv(fields, "phuong_phap_tai_chinh"),
        field_val(data.get("chuong_iii") or {}, "muc_4_tai_chinh"),
        str((data.get("decisions") or {}).get("Q12_phuong_phap_tai_chinh") or ""),
    ]
    blob = " ".join(chunks).lower()
    if "đánh giá" in blob or "danh gia" in blob:
        return "gia_danh_gia"
    if "thấp nhất" in blob or "thap nhat" in blob:
        return "gia_thap_nhat"
    return ""


def build_ecdnt_28_2(pp_tc: str) -> str:
    """Chọn đúng 1 nhánh — không giữ 2 dòng hướng dẫn."""
    mo_dau = (
        "Cách tính ưu đãi: __\n"
        "Sản phẩm đổi mới sáng tạo quy định tại khoản 4 Điều 6 của Nghị định số "
        "214/2025/NĐ-CP là dịch vụ phi tư vấn của doanh nghiệp khởi nghiệp sáng tạo "
        "trong nước được hưởng ưu đãi như sau:"
    )
    if pp_tc == "gia_thap_nhat":
        bullet = (
            "- Nhà thầu không thuộc đối tượng được hưởng ưu đãi phải cộng thêm một khoản "
            "tiền bằng 7,5% giá dự thầu trừ đi giá trị giảm giá (nếu có) vào giá dự thầu "
            "trừ đi giá trị giảm giá (nếu có) của nhà thầu đó để so sánh, xếp hạng"
        )
    elif pp_tc == "gia_danh_gia":
        bullet = (
            "- Nhà thầu không thuộc đối tượng được hưởng ưu đãi phải cộng thêm một khoản "
            "tiền bằng 7,5% giá dự thầu trừ đi giá trị giảm giá (nếu có) vào giá đánh giá "
            "của nhà thầu đó để so sánh, xếp hạng"
        )
    else:
        return ""
    return f"{mo_dau}\n{bullet}"


def build_ecdnt_31_4(pp_tc: str, original: str) -> str:
    """Chọn 1 câu trích; giữ đoạn nhiều phần (lô)."""
    if pp_tc == "gia_thap_nhat":
        chosen = (
            "- Nhà thầu có giá dự thầu sau khi trừ đi giá trị giảm giá (nếu có) thấp nhất"
        )
    elif pp_tc == "gia_danh_gia":
        chosen = "- Nhà thầu có giá đánh giá thấp nhất"
    else:
        return ""
    # Giữ đoạn lô từ template gốc
    lo = ""
    m = re.search(r"(Đối với gói thầu được chia làm nhiều phần.*)", original, flags=re.S)
    if m:
        lo = m.group(1).strip()
    else:
        lo = (
            "Đối với gói thầu được chia làm nhiều phần (lô), việc đánh giá E-HSDT và xét "
            "duyệt trúng thầu sẽ được thực hiện trên cơ sở bảo đảm: giá đề nghị trúng thầu "
            "của cả gói thầu không vượt giá gói thầu được duyệt mà không so sánh với ước "
            "tính chi phí của từng phần."
        )
    return f"{chosen}\n{lo}"


def build_ecdnt_31_5(loai: str) -> str:
    """loai: tron_goi_hoac_don_gia | ket_qua_dau_ra — theo nhánh template."""
    prefix = "Giá gói thầu để làm căn cứ xét duyệt trúng thầu: "
    if loai in ("tron_goi", "don_gia", "thoi_gian", "tron_goi_hoac_don_gia"):
        return (
            prefix
            + "giá gói thầu được phê duyệt trong kế hoạch lựa chọn nhà thầu. Trường hợp "
            "dự toán của gói thầu được phê duyệt thấp hơn hoặc cao hơn giá gói thầu được "
            "phê duyệt thì dự toán này sẽ thay thế giá gói thầu để làm cơ sở xét duyệt "
            "trúng thầu."
        )
    if loai in ("ket_qua_dau_ra", "ty_le"):
        return (
            prefix
            + "giá gói thầu không bao gồm chi phí dự phòng phát sinh khối lượng và chi phí "
            "dự phòng trượt giá có thể xảy ra trong quá trình thực hiện hợp đồng. Trường "
            "hợp dự toán của gói thầu được phê duyệt thấp hơn hoặc cao hơn giá gói thầu "
            "được phê duyệt thì dự toán này (không bao gồm chi phí dự phòng phát sinh "
            "khối lượng và chi phí dự phòng trượt giá) sẽ thay thế giá gói thầu để làm cơ "
            "sở xét duyệt trúng thầu."
        )
    return ""


# --- Chương II ---
def fill_chuong_ii(doc: Document, data: dict) -> None:
    fields = ((data.get("chuong_ii") or {}).get("fields")) or {}
    if not fields:
        return

    table = None
    for t in doc.tables:
        if len(t.columns) >= 2 and t.rows[0].cells[0].text.strip().startswith("E-CDNT"):
            table = t
            break
    if table is None:
        print("Cảnh báo: không tìm thấy bảng E-CDNT (Ch II)", file=sys.stderr)
        return

    iii = data.get("chuong_iii") or {}
    pp = iii.get("phuong_phap_de_xuat") or {}
    method = pp.get("value") if isinstance(pp, dict) else str(pp or "")
    tc = iii.get("muc_4_tai_chinh") or {}
    fin = tc.get("value") if isinstance(tc, dict) else ""

    skipped = []

    for row in table.rows:
        code = ecdnt_code(row.cells[0].text)
        if not code:
            continue
        cell = row.cells[1]
        original = cell.text

        # --- chỉ các mã được phép điền; match chính xác ---
        if code == "1.1":
            if has_value(fields, "chu_dau_tu"):
                set_cell_text(cell, f"Tên Chủ đầu tư: {fv(fields, 'chu_dau_tu')}")
            else:
                skipped.append(code)

        elif code == "1.2":
            if has_value(fields, "ten_goi_thau") or has_value(fields, "ten_du_an"):
                set_cell_text(
                    cell,
                    f"Tên gói thầu: {fv(fields, 'ten_goi_thau')}. "
                    f"Tên dự án/dự toán mua sắm: {fv(fields, 'ten_du_an')}.",
                )
            else:
                skipped.append(code)

        elif code == "3":
            if has_value(fields, "nguon_von"):
                set_cell_text(cell, f"Nguồn vốn: {fv(fields, 'nguon_von')}")
            else:
                skipped.append(code)

        elif code == "5.1":
            # Bắt buộc chuỗi đầy đủ tên + địa chỉ (Decision). Không tự lấy tên ngắn.
            if has_value(fields, "chu_dau_tu_day_du"):
                cdt = fv(fields, "chu_dau_tu_day_du")
                new = re.sub(
                    r"với:\s*_{0,}(\[.*?\]|___+)?",
                    f"với: {cdt}",
                    original,
                    count=1,
                    flags=re.I | re.S,
                )
                if new == original:
                    new = replace_blank(original, r"_{3,}|\[.*?\]", cdt)
                set_cell_text(cell, new)
            else:
                skipped.append(
                    "5.1(cần chu_dau_tu_day_du = tên+địa chỉ; không dùng tên ngắn)"
                )

        elif code == "7.5":
            # Bắt buộc Decision; không suy từ thuyết trình RFP
            if has_value(fields, "hoi_nghi_tien_dau_thau"):
                set_cell_text(
                    cell, f"Hội nghị tiền đấu thầu: {fv(fields, 'hoi_nghi_tien_dau_thau')}"
                )
            else:
                skipped.append(code)

        elif code == "10.8":
            if has_value(fields, "tai_lieu_nop_kem"):
                set_cell_text(
                    cell,
                    "Nhà thầu phải nộp cùng với E-HSDT các tài liệu sau đây:\n"
                    + fv(fields, "tai_lieu_nop_kem"),
                )
            else:
                skipped.append(code)

        elif code == "12.1":
            if has_value(fields, "phuong_an_thay_the"):
                pa = fv(fields, "phuong_an_thay_the")
                low = pa.lower()
                if "không" in low or "khong" in low:
                    set_cell_text(
                        cell,
                        "Nhà thầu không được phép nộp đề xuất phương án kỹ thuật thay thế.",
                    )
                else:
                    set_cell_text(
                        cell,
                        f"Nhà thầu được phép nộp đề xuất phương án kỹ thuật thay thế. ({pa})",
                    )
            else:
                skipped.append(code)

        elif code == "17.1":
            if has_value(fields, "hieu_luc_hsdt"):
                set_cell_text(
                    cell,
                    f"Thời hạn hiệu lực của E-HSDT là: {fv(fields, 'hieu_luc_hsdt')}, "
                    "kể từ ngày có thời điểm đóng thầu.",
                )
            else:
                skipped.append(code)

        elif code == "18.2":
            # Giữ khung pháp lý; chỉ điền giá trị nếu có và đang ÁP DỤNG bảo đảm
            bao_dam_flag = fv(fields, "bao_dam_du_thau_ap_dung").lower()
            if "không áp dụng" in bao_dam_flag or "khong ap dung" in bao_dam_flag:
                # Không đụng 18.2 khi chọn không áp dụng ở 18.9
                skipped.append(f"{code}(giữ template vì 18.9 Không áp dụng)")
            elif has_value(fields, "bao_dam_du_thau_gia_tri"):
                val = fv(fields, "bao_dam_du_thau_gia_tri")
                new = re.sub(
                    r"Giá trị bảo đảm dự thầu:\s*_{0,}(\[.*?\]|___+)?",
                    f"Giá trị bảo đảm dự thầu: {val}",
                    original,
                    count=1,
                    flags=re.I | re.S,
                )
                set_cell_text(cell, new)
            else:
                skipped.append(code)

        elif code == "18.4":
            if has_value(fields, "bao_dam_hoan_tra_ngay"):
                days = fv(fields, "bao_dam_hoan_tra_ngay")
                set_cell_text(
                    cell,
                    f"Thời gian hoàn trả hoặc giải tỏa bảo đảm dự thầu đối với nhà thầu "
                    f"không được lựa chọn: {days} ngày, kể từ ngày kết quả lựa chọn nhà thầu "
                    "được phê duyệt.",
                )
            else:
                skipped.append(code)

        elif code == "18.9":
            if has_value(fields, "bao_dam_du_thau_ap_dung"):
                set_cell_text(
                    cell, f"Bảo đảm dự thầu: {fv(fields, 'bao_dam_du_thau_ap_dung')}"
                )
            else:
                skipped.append(code)

        elif code == "27.2":
            if has_value(fields, "nha_thau_phu"):
                set_cell_text(
                    cell, f"Giá trị tối đa dành cho nhà thầu phụ: {fv(fields, 'nha_thau_phu')}"
                )
            else:
                skipped.append(code)

        elif code == "28.2":
            # Không được giữ 2 dòng "Chọn phương pháp…". Chọn đúng 1 nhánh.
            if has_value(fields, "uu_dai_cach_tinh"):
                set_cell_text(cell, fv(fields, "uu_dai_cach_tinh"))
            else:
                built = build_ecdnt_28_2(detect_phuong_phap_tai_chinh(data, fields))
                if built:
                    set_cell_text(cell, built)
                else:
                    skipped.append("28.2(chưa có phương pháp TC để chọn 1 nhánh)")

        elif code == "28.8":
            if has_value(fields, "uu_dai_khac"):
                set_cell_text(
                    cell, f"Các ưu đãi khác (nếu có): {fv(fields, 'uu_dai_khac')}"
                )
            else:
                skipped.append(code)

        elif code == "29.1":
            # Văn phong giống template (gạch đầu dòng), không nhét Decision code
            kt = ""
            if method:
                if "cham" in method.lower() or "chấm" in method.lower():
                    kt = "Chấm điểm"
                elif "dat" in method.lower() or "đạt" in method.lower():
                    kt = "Đạt/Không đạt"
                else:
                    kt = method
            tc_txt = ""
            pp_tc = detect_phuong_phap_tai_chinh(data, fields)
            if pp_tc == "gia_thap_nhat":
                tc_txt = "phương pháp giá thấp nhất"
            elif pp_tc == "gia_danh_gia":
                tc_txt = "phương pháp giá đánh giá"
            elif fin:
                tc_txt = fin
            if kt or tc_txt:
                lines = [
                    "Phương pháp đánh giá E-HSDT là:",
                    "- Đánh giá về năng lực và kinh nghiệm: Đạt/Không đạt.",
                    f"- Đánh giá về kỹ thuật: {kt or '____'}.",
                    f"- Đánh giá về tài chính: {tc_txt or '____'}.",
                ]
                set_cell_text(cell, "\n".join(lines))
            else:
                skipped.append(code)

        elif code == "29.3":
            if has_value(fields, "xep_hang_nha_thau"):
                set_cell_text(cell, fv(fields, "xep_hang_nha_thau"))
            else:
                pp_tc = detect_phuong_phap_tai_chinh(data, fields)
                if pp_tc == "gia_thap_nhat":
                    set_cell_text(
                        cell,
                        "Xếp hạng nhà thầu: Nhà thầu có giá dự thầu sau khi trừ đi giá trị "
                        "giảm giá (nếu có) thấp nhất được xếp hạng thứ nhất.",
                    )
                elif pp_tc == "gia_danh_gia":
                    set_cell_text(
                        cell,
                        "Xếp hạng nhà thầu: Nhà thầu có giá đánh giá thấp nhất được xếp "
                        "hạng thứ nhất.",
                    )
                else:
                    skipped.append(code)

        elif code == "31.4":
            if has_value(fields, "ecdnt_31_4"):
                set_cell_text(cell, fv(fields, "ecdnt_31_4"))
            else:
                built = build_ecdnt_31_4(
                    detect_phuong_phap_tai_chinh(data, fields), original
                )
                if built:
                    set_cell_text(cell, built)
                else:
                    skipped.append("31.4(chưa có phương pháp TC để chọn 1 nhánh)")

        elif code == "31.5":
            if has_value(fields, "ecdnt_31_5"):
                set_cell_text(cell, fv(fields, "ecdnt_31_5"))
            elif has_value(fields, "loai_hop_dong"):
                built = build_ecdnt_31_5(fv(fields, "loai_hop_dong"))
                if built:
                    set_cell_text(cell, built)
                else:
                    skipped.append("31.5(loai_hop_dong không nhận diện được)")
            else:
                skipped.append("31.5(cần hỏi loại hợp đồng / ecdnt_31_5)")

        elif code == "34.1":
            if has_value(fields, "ty_le_tang_giam_khoi_luong"):
                set_cell_text(cell, fv(fields, "ty_le_tang_giam_khoi_luong"))
            else:
                skipped.append(code)

        elif code == "34.2":
            if has_value(fields, "tuy_chon_mua_them"):
                set_cell_text(cell, fv(fields, "tuy_chon_mua_them"))
            else:
                skipped.append(code)

        elif code == "38.1":
            if has_value(fields, "nguoi_co_tham_quyen"):
                set_cell_text(cell, fv(fields, "nguoi_co_tham_quyen"))
            else:
                skipped.append(code)

        elif code == "38.2":
            if has_value(fields, "quy_trinh_kien_nghi"):
                set_cell_text(cell, fv(fields, "quy_trinh_kien_nghi"))
            else:
                skipped.append(code)

        # Các mã khác: không đụng

    if skipped:
        print(
            "Ch II bỏ qua (thiếu Decision/Fact, giữ template): " + ", ".join(skipped),
            file=sys.stderr,
        )


def fill_bang_nang_luc(doc: Document, muc2: dict) -> None:
    """Điền Bảng số 01 — tiêu chuẩn năng lực & kinh nghiệm (không bỏ quên)."""
    bang = muc2.get("bang_tieu_chuan") or {}
    if not bang:
        print(
            "Cảnh báo: thiếu chuong_iii.muc_2_nang_luc.bang_tieu_chuan — "
            "không điền bảng năng lực (cần hỏi Manual).",
            file=sys.stderr,
        )
        return

    table = None
    for t in doc.tables:
        if not t.rows:
            continue
        head = t.rows[0].cells[0].text
        if "NĂNG LỰC" in head.upper() and "KINH NGHIỆM" in head.upper():
            table = t
            break
    if table is None or len(table.rows) < 8:
        print("Cảnh báo: không tìm thấy Bảng số 01 năng lực", file=sys.stderr)
        return

    # Hàng dữ liệu: row 4=TT1, 5=TT2, 6=TT3, 7=TT4 — cột Yêu cầu = cells[2]
    # Field có thể là string hoặc {value, confidence, evidence,...} → chỉ lấy value
    # TT1: thay năm ___
    nam = as_plain_text(bang.get("lich_su_nam"))
    if nam:
        yeu = table.rows[4].cells[2].text
        # Template dạng: năm ___(1) → năm 2024 (bỏ chỗ trống, giữ hoặc bỏ footnote tùy mẫu)
        yeu2 = re.sub(
            r"năm\s*_{2,}\s*(\(\d+\))?",
            f"năm {nam}",
            yeu,
            count=1,
        )
        if yeu2 == yeu:
            yeu2 = re.sub(r"_{2,}\s*(\(\d+\))?", nam, yeu, count=1)
        set_cell_text(table.rows[4].cells[2], yeu2)

    # TT3 doanh thu: "Không áp dụng" hoặc văn bản số tiền
    dt = as_plain_text(bang.get("doanh_thu_yeu_cau"))
    if dt:
        set_cell_text(table.rows[6].cells[2], dt)

    # TT4 hợp đồng tương tự
    hd = as_plain_text(bang.get("hop_dong_tuong_tu_yeu_cau"))
    if hd:
        set_cell_text(table.rows[7].cells[2], hd)


# --- Chương III ---
def fill_chuong_iii(doc: Document, data: dict) -> None:
    iii = data.get("chuong_iii") or {}
    muc2 = iii.get("muc_2_nang_luc") or {}
    fill_bang_nang_luc(doc, muc2)

    nhan_su = muc2.get("nhan_su_chu_chot") or {}
    roles_detail = nhan_su.get("chi_tiet") or []
    # Fallback cũ: chỉ có tên role
    if not roles_detail:
        roles_detail = [
            {
                "vi_tri": r,
                "so_luong": "01",
                "kinh_nghiem": "",
                "chung_chi": "Không yêu cầu",
            }
            for r in (nhan_su.get("roles_toi_thieu") or [])
        ]

    ns_table = None
    for t in doc.tables:
        if len(t.columns) >= 4 and "Vị trí công việc" in t.rows[0].cells[1].text:
            ns_table = t
            break
    if ns_table is not None and roles_detail:
        ensure_rows(ns_table, len(roles_detail), header_rows=1)
        for i, role in enumerate(roles_detail):
            row = ns_table.rows[1 + i]
            set_cell_text(row.cells[0], str(role.get("stt") or (i + 1)))
            set_cell_text(row.cells[1], str(role.get("vi_tri") or ""))
            set_cell_text(row.cells[2], str(role.get("so_luong") or "01"))
            kn = str(role.get("kinh_nghiem") or "").strip()
            set_cell_text(
                row.cells[3],
                kn if kn else "Chưa chốt (chờ Decision Manual)",
            )
            if len(row.cells) > 4:
                cc = str(role.get("chung_chi") or "").strip()
                set_cell_text(row.cells[4], cc if cc else "Không yêu cầu")

    bang = iii.get("bang_cham_diem_proposed") or iii.get("bang_cham_diem") or []
    if not bang:
        return

    anchor = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Mục 3") and "kỹ thuật" in t.lower():
            anchor = p
    if anchor is None:
        for p in doc.paragraphs:
            if "Tiêu chuẩn đánh giá về kỹ thuật" in p.text:
                anchor = p
                break
    if anchor is None:
        print("Cảnh báo: không tìm thấy chỗ chèn bảng chấm điểm Ch III", file=sys.stderr)
        return

    pp = iii.get("phuong_phap_de_xuat") or {}
    san_nhom = pp.get("san_nhom") or {}
    intro_lines = [
        f"Bảng tiêu chuẩn đánh giá về kỹ thuật (thang điểm {pp.get('tong_diem', 100)}).",
        f"Sàn tổng: {pp.get('san_tong', 70)}."
        + (
            f" Sàn nhóm: A≥{san_nhom.get('A')}, B≥{san_nhom.get('B')}, C≥{san_nhom.get('C')}."
            if san_nhom
            else ""
        ),
        "Cách chấm từng tiêu chí (trên cơ sở nội dung trình bày trong E-HSDT): "
        "Không nêu/không đáp ứng = 0; nêu sơ sài = điểm tối thiểu; đủ và rõ = điểm tối đa "
        "(chi tiết theo cột Thang điểm chi tiết).",
    ]
    cursor = make_paragraph_after(anchor, "\n".join(intro_lines))

    nrows = 1 + len(bang) + 1
    table = doc.add_table(rows=nrows, cols=5)
    set_table_grid(table)
    headers = [
        "STT",
        "Nội dung đánh giá",
        "Điểm tối đa",
        "Điểm tối thiểu",
        "Thang điểm chi tiết",
    ]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)

    total_max = 0
    for i, item in enumerate(bang):
        r = table.rows[1 + i]
        set_cell_text(r.cells[0], str(item.get("stt", "")))
        set_cell_text(r.cells[1], str(item.get("noi_dung", "")))
        dm = item.get("diem_toi_da")
        set_cell_text(r.cells[2], "" if dm is None else str(dm))
        dmin = item.get("diem_toi_thieu")
        set_cell_text(r.cells[3], "" if dmin is None else str(dmin))
        set_cell_text(r.cells[4], str(item.get("thang_chi_tiet", "")))
        # Chỉ cộng điểm tiêu chí con (A1…) không cộng hàng nhóm A/B/C nếu đã có sẵn tổng
        stt = str(item.get("stt", ""))
        if isinstance(dm, (int, float)) and re.match(r"^[A-Z]\d", stt):
            total_max += dm
        elif isinstance(dm, (int, float)) and stt in ("A", "B", "C"):
            pass
        elif isinstance(dm, (int, float)) and not re.match(r"^[A-Z]$", stt):
            total_max += dm

    last = table.rows[-1]
    set_cell_text(last.cells[0], "", bold=True)
    set_cell_text(last.cells[1], "TỔNG CỘNG", bold=True)
    set_cell_text(last.cells[2], str(pp.get("tong_diem") or total_max or ""), bold=True)
    set_cell_text(last.cells[3], str(pp.get("san_tong", "")), bold=True)
    san_txt = "Đạt KT (trên HSDT) khi tổng ≥ sàn"
    if san_nhom:
        san_txt += (
            f" và A≥{san_nhom.get('A')}, B≥{san_nhom.get('B')}, C≥{san_nhom.get('C')}"
        )
    set_cell_text(last.cells[4], san_txt)

    tbl = table._tbl
    tbl.getparent().remove(tbl)
    cursor._p.addnext(tbl)


# --- Chương IV ---
def fill_chuong_iv(doc: Document, data: dict) -> None:
    hang = (data.get("chuong_iv") or {}).get("hang_muc") or []
    if not hang:
        return
    table = None
    for t in doc.tables:
        if len(t.columns) >= 7 and "Danh mục dịch vụ" in t.rows[0].cells[1].text:
            if len(t.rows) <= 10:
                table = t
                break
            if table is None:
                table = t
    if table is None:
        print("Cảnh báo: không tìm thấy bảng phạm vi Ch IV", file=sys.stderr)
        return

    header_rows = 2 if "(1)" in table.rows[1].cells[0].text else 1
    ensure_rows(table, len(hang), header_rows=header_rows)
    for i, item in enumerate(hang):
        row = table.rows[header_rows + i]
        dia = str(item.get("dia_diem") or "").strip()
        if not dia or dia.upper().startswith("TBD"):
            dia = "Chưa chốt (chờ Decision Manual)"
        vals = [
            str(item.get("stt", "")),
            str(item.get("ten", "")),
            str(item.get("mo_ta", "")),
            str(item.get("khoi_luong", "")),
            str(item.get("don_vi", "")),
            dia,
            str(item.get("ngay_hoan_thanh", "")),
        ]
        for ci, val in enumerate(vals):
            if ci < len(row.cells):
                set_cell_text(row.cells[ci], val)


# --- Chương V ---
def fill_chuong_v(doc: Document, data: dict) -> None:
    cv = data.get("chuong_v") or {}
    sec = cv.get("sections") or {}
    fields_ii = ((data.get("chuong_ii") or {}).get("fields")) or {}
    nfr = cv.get("nfr_table") or []
    chuc = cv.get("chuc_nang_table") or []

    p_v = find_para(doc, "Chương V. YÊU CẦU") or find_para(doc, "Chương V. Yêu cầu")
    if p_v is None:
        raise RuntimeError("Không tìm thấy Chương V trong template")

    p1 = find_para(doc, "1. Giới thiệu", p_v)
    p2 = find_para(doc, "2. Mục tiêu", p_v)
    p3 = find_para(doc, "3. Yêu cầu kỹ thuật", p_v)
    p4 = find_para(doc, "4. Giải pháp", p_v)
    p5 = find_para(doc, "5. Quy định về kiểm tra", p_v)
    p_end = find_para(doc, "Phần 3.", p_v) or find_para(doc, "Chương VI.", p_v)
    if None in (p1, p2, p3, p4, p5, p_end):
        raise RuntimeError("Thiếu mục Ch V trong template")

    delete_between(p1, p2)
    delete_between(p2, p3)
    delete_between(p3, p4)
    delete_between(p4, p5)
    delete_between(p5, p_end)

    p1 = find_para(doc, "1. Giới thiệu", p_v)
    p2 = find_para(doc, "2. Mục tiêu", p_v)
    p3 = find_para(doc, "3. Yêu cầu kỹ thuật", p_v)
    p4 = find_para(doc, "4. Giải pháp", p_v)
    p5 = find_para(doc, "5. Quy định về kiểm tra", p_v)

    p = p1
    p = make_paragraph_after(p, f"Tên dự án: {fv(fields_ii, 'ten_du_an')}")
    p = make_paragraph_after(p, f"Tên gói thầu: {fv(fields_ii, 'ten_goi_thau')}")
    p = make_paragraph_after(p, f"Chủ đầu tư: {fv(fields_ii, 'chu_dau_tu')}")
    if field_val(sec, "gioi_thieu"):
    make_paragraph_after(p, field_val(sec, "gioi_thieu"))

    if field_val(sec, "muc_tieu"):
        make_paragraph_after(p2, field_val(sec, "muc_tieu"))

    cursor = make_paragraph_after(p3, "3.1. Yêu cầu phi chức năng", bold=True)
    if nfr:
        t = doc.add_table(rows=1 + len(nfr), cols=3)
        set_table_grid(t)
        for i, h in enumerate(["STT", "Nội dung", "Yêu cầu"]):
            set_cell_text(t.rows[0].cells[i], h, bold=True)
        for i, item in enumerate(nfr):
            set_cell_text(t.rows[1 + i].cells[0], str(item.get("stt", "")))
            set_cell_text(t.rows[1 + i].cells[1], str(item.get("noi_dung", "")))
            set_cell_text(t.rows[1 + i].cells[2], str(item.get("yeu_cau", "")))
        tbl = t._tbl
        tbl.getparent().remove(tbl)
        cursor._p.addnext(tbl)
        new_p = OxmlElement("w:p")
        tbl.addnext(new_p)
        cursor = Paragraph(new_p, p3._parent)
    elif field_val(sec, "nfr"):
        cursor = make_paragraph_after(cursor, field_val(sec, "nfr"))

    cursor = make_paragraph_after(cursor, "3.2. Yêu cầu chức năng theo phân hệ", bold=True)
    if chuc:
        t = doc.add_table(rows=1 + len(chuc), cols=3)
        set_table_grid(t)
        for i, h in enumerate(["STT", "Nội dung công việc", "Yêu cầu"]):
            set_cell_text(t.rows[0].cells[i], h, bold=True)
        for i, item in enumerate(chuc):
            set_cell_text(t.rows[1 + i].cells[0], str(item.get("stt", "")))
            set_cell_text(t.rows[1 + i].cells[1], str(item.get("noi_dung", "")))
            set_cell_text(t.rows[1 + i].cells[2], str(item.get("yeu_cau", "")))
        tbl = t._tbl
        tbl.getparent().remove(tbl)
        cursor._p.addnext(tbl)
    elif field_val(sec, "chuc_nang"):
        make_paragraph_after(cursor, field_val(sec, "chuc_nang"))

    if field_val(sec, "giai_phap_yeu_cau_nha_thau"):
        make_paragraph_after(p4, field_val(sec, "giai_phap_yeu_cau_nha_thau"))
    if field_val(sec, "ban_giao_nghiem_thu"):
        make_paragraph_after(p5, field_val(sec, "ban_giao_nghiem_thu"))


def main() -> int:
    parser = argparse.ArgumentParser(description="Fill HSMT từ 01_hsmt.json")
    parser.add_argument(
        "--force",
        action="store_true",
        help="Điền cả khi status chưa approved (chỉ để thử)",
    )
    args = parser.parse_args()

    if not JSON_PATH.exists():
        print(f"Thiếu {JSON_PATH}", file=sys.stderr)
        return 1
    data = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    if data.get("status") != "approved" and not args.force:
        print(
            f"Từ chối fill: status={data.get('status')!r}. "
            "Chỉ chạy khi status=approved (hoặc thêm --force).",
            file=sys.stderr,
        )
        return 2

    template = find_template(data)
    shutil.copy2(template, OUT_PATH)
    doc = Document(str(OUT_PATH))
    fill_chuong_ii(doc, data)
    fill_chuong_iii(doc, data)
    fill_chuong_iv(doc, data)
    fill_chuong_v(doc, data)
    doc.save(str(OUT_PATH))
    print(f"Đã ghi {OUT_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
