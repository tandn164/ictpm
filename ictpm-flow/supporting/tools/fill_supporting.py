#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fill template supporting từ 01_supporting.json (approved).

  python3 ictpm-flow/supporting/tools/fill_supporting.py --list
  python3 ictpm-flow/supporting/tools/fill_supporting.py --doc stakeholder_analysis
  python3 ictpm-flow/supporting/tools/fill_supporting.py --all
  python3 ictpm-flow/supporting/tools/fill_supporting.py --all --force
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

MODULE = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = MODULE / "templates"
OUT_DIR = MODULE / "artifacts"
JSON_PATH = OUT_DIR / "01_supporting.json"

PLACEHOLDER_RE = re.compile(r"\{\{([a-z0-9_]+)\}\}")
FONT = Font(name="Times New Roman", size=14)
THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

CATALOG = {
    "stakeholder_analysis": ("01_stakeholder_analysis.xlsx", "stakeholder_analysis_draft.xlsx"),
    "issues_actions": ("02_issues_actions.xlsx", "issues_actions_draft.xlsx"),
    "phase_exit_checklist": ("03_phase_exit_checklist.docx", "phase_exit_checklist_draft.docx"),
    "peer_evaluation": ("04_peer_evaluation.docx", "peer_evaluation_draft.docx"),
    "selection_matrix": ("05_selection_matrix.xlsx", "selection_matrix_draft.xlsx"),
    "portfolio": ("06_portfolio.xlsx", "portfolio_draft.xlsx"),
    "stakeholder_post_eval": (
        "07_stakeholder_post_evaluation.docx",
        "stakeholder_post_evaluation_draft.docx",
    ),
    "post_impl_survey": (
        "08_post_implementation_survey.docx",
        "post_implementation_survey_draft.docx",
    ),
}


def plain(cur, default: str = "") -> str:
    if cur is None:
        return default
    if isinstance(cur, dict):
        v = cur.get("value")
        if v is None or (isinstance(v, str) and not str(v).strip()):
            v = cur.get("proposed")
        if v is None:
            return default
        if isinstance(v, (list, tuple)):
            return ", ".join(str(x) for x in v)
        return str(v).strip()
    if isinstance(cur, (list, tuple)):
        return ", ".join(str(x) for x in cur)
    return str(cur).strip()


def load_data() -> dict:
    if not JSON_PATH.exists():
        raise SystemExit(f"Thiếu {JSON_PATH}")
    return json.loads(JSON_PATH.read_text(encoding="utf-8"))


def meta_from(data: dict) -> dict[str, str]:
    proj = data.get("project") or {}
    docs = data.get("documents") or {}
    peer = docs.get("peer_evaluation") or {}
    post = docs.get("stakeholder_post_eval") or {}
    survey = docs.get("post_impl_survey") or {}
    phase = docs.get("phase_exit_checklist") or {}
    sel = docs.get("selection_matrix") or {}
    return {
        "project_name": plain(proj.get("name")),
        "project_code": plain(proj.get("code")),
        "company": plain(proj.get("company")),
        "pm": plain(proj.get("pm")),
        "sponsor": plain(proj.get("sponsor")),
        "date": plain(proj.get("date")),
        "phase": plain(phase.get("phase"), "—"),
        "gate_decision": plain(phase.get("gate_decision")),
        "gate_conditions": plain(phase.get("gate_conditions")),
        "evaluator": plain(peer.get("evaluator")),
        "team_strengths": plain(peer.get("team_strengths")),
        "team_improve": plain(peer.get("team_improve")),
        "personal_lesson": plain(peer.get("personal_lesson")),
        "stakeholder_name": plain(post.get("stakeholder_name")),
        "what_went_well": plain(post.get("what_went_well")),
        "main_challenges": plain(post.get("main_challenges")),
        "recommendations": plain(post.get("recommendations")),
        "rating_objectives": plain(post.get("rating_objectives")),
        "rating_communication": plain(post.get("rating_communication")),
        "rating_change_risk": plain(post.get("rating_change_risk")),
        "rating_ops": plain(post.get("rating_ops")),
        "respondent": plain(survey.get("respondent")),
        "role": plain(survey.get("role")),
        "improve_now": plain(survey.get("improve_now")),
        "best_practice": plain(survey.get("best_practice")),
        "selection_recommendation": plain(sel.get("recommendation")),
    }


def replace_text(text: str, meta: dict[str, str]) -> str:
    def repl(m: re.Match) -> str:
        return meta.get(m.group(1), "")

    return PLACEHOLDER_RE.sub(repl, text)


def style_xlsx_cell(cell) -> None:
    cell.font = FONT
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = THIN


def _docx_font(run, size: int = 14, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")


def set_para_text(p, text: str, *, bold: bool = False) -> None:
    text = "" if text is None else str(text)
    for r in p.runs:
        r.text = ""
    if p.runs:
        p.runs[0].text = text
        _docx_font(p.runs[0], bold=bold)
    else:
        run = p.add_run(text)
        _docx_font(run, bold=bold)


def replace_in_docx(doc: Document, meta: dict[str, str]) -> None:
    for p in doc.paragraphs:
        if "{{" in p.text:
            set_para_text(p, replace_text(p.text, meta))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if "{{" in p.text:
                        set_para_text(p, replace_text(p.text, meta))


def fill_stakeholder(src: Path, dest: Path, data: dict, meta: dict[str, str]) -> None:
    shutil.copy2(src, dest)
    wb = load_workbook(dest)
    ws = wb["Stakeholder Register"]
    for row in ws.iter_rows(min_row=1, max_row=3):
        for cell in row:
            if isinstance(cell.value, str) and "{{" in cell.value:
                cell.value = replace_text(cell.value, meta)
                cell.font = FONT
    rows = (data.get("documents") or {}).get("stakeholder_analysis", {}).get("stakeholders") or []
    for i, st in enumerate(rows):
        r = 5 + i
        vals = [
            i + 1,
            st.get("name", ""),
            st.get("role", ""),
            st.get("power", ""),
            st.get("interest", ""),
            st.get("strategy", ""),
            st.get("owner", ""),
            st.get("notes") or st.get("engagement", ""),
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            style_xlsx_cell(cell)

    ws2 = wb["Ma trận Quyền lực-Quan tâm"]
    for row in ws2.iter_rows(min_row=1, max_row=3):
        for cell in row:
            if isinstance(cell.value, str) and "{{" in cell.value:
                cell.value = replace_text(cell.value, meta)
                cell.font = FONT
    buckets = {
        ("High", "Low"): [],
        ("High", "Middle"): [],
        ("High", "High"): [],
        ("Middle", "Low"): [],
        ("Middle", "Middle"): [],
        ("Middle", "High"): [],
        ("Low", "Low"): [],
        ("Low", "Middle"): [],
        ("Low", "High"): [],
    }
    for st in rows:
        key = (st.get("power", "Middle"), st.get("interest", "Middle"))
        if key not in buckets:
            key = ("Middle", "Middle")
        buckets[key].append(st.get("name", ""))
    # rows 5-7 = High/Middle/Low power; cols 2-4 = Low/Middle/High interest
    power_order = ["High", "Middle", "Low"]
    interest_order = ["Low", "Middle", "High"]
    for ri, pwr in enumerate(power_order):
        for ci, it in enumerate(interest_order):
            text = "\n".join(buckets[(pwr, it)])
            cell = ws2.cell(5 + ri, 2 + ci, text)
            style_xlsx_cell(cell)
    wb.save(dest)


def fill_issues(src: Path, dest: Path, data: dict, meta: dict[str, str]) -> None:
    shutil.copy2(src, dest)
    wb = load_workbook(dest)
    ws = wb.active
    for row in ws.iter_rows(min_row=1, max_row=3):
        for cell in row:
            if isinstance(cell.value, str) and "{{" in cell.value:
                cell.value = replace_text(cell.value, meta)
                cell.font = FONT
    items = (data.get("documents") or {}).get("issues_actions", {}).get("items") or []
    for i, it in enumerate(items):
        r = 5 + i
        vals = [
            it.get("id", f"IA-{i+1:03d}"),
            it.get("type", "Issue"),
            it.get("description", ""),
            it.get("priority", ""),
            it.get("status", ""),
            it.get("owner", ""),
            it.get("due", ""),
            it.get("action", ""),
            it.get("related_risk", ""),
        ]
        for c, v in enumerate(vals, 1):
            style_xlsx_cell(ws.cell(r, c, v))
    wb.save(dest)


def fill_selection(src: Path, dest: Path, data: dict, meta: dict[str, str]) -> None:
    shutil.copy2(src, dest)
    wb = load_workbook(dest)
    ws = wb.active
    for row in ws.iter_rows(min_row=1, max_row=3):
        for cell in row:
            if isinstance(cell.value, str) and "{{" in cell.value:
                cell.value = replace_text(cell.value, meta)
                cell.font = FONT
    criteria = (data.get("documents") or {}).get("selection_matrix", {}).get("criteria") or []
    for i, cr in enumerate(criteria):
        r = 5 + i
        vals = [
            cr.get("name", ""),
            cr.get("weight", ""),
            cr.get("option_a", ""),
            cr.get("option_b", ""),
            cr.get("option_c", ""),
            cr.get("notes", ""),
        ]
        for c, v in enumerate(vals, 1):
            style_xlsx_cell(ws.cell(r, c, v))
    totals = (data.get("documents") or {}).get("selection_matrix", {}).get("totals") or {}
    if totals:
        ws.cell(16, 1, "Tổng điểm có trọng số")
        style_xlsx_cell(ws.cell(16, 1))
        ws.cell(16, 1).font = Font(name="Times New Roman", size=14, bold=True)
        for c, key in enumerate(["", "weight_sum", "option_a", "option_b", "option_c"], 1):
            if c == 1:
                continue
            style_xlsx_cell(ws.cell(16, c, totals.get(key, "")))
    wb.save(dest)


def fill_portfolio(src: Path, dest: Path, data: dict, meta: dict[str, str]) -> None:
    shutil.copy2(src, dest)
    wb = load_workbook(dest)
    ws = wb.active
    for row in ws.iter_rows(min_row=1, max_row=3):
        for cell in row:
            if isinstance(cell.value, str) and "{{" in cell.value:
                cell.value = replace_text(cell.value, meta)
                cell.font = FONT
    projects = (data.get("documents") or {}).get("portfolio", {}).get("projects") or []
    for i, p in enumerate(projects):
        r = 5 + i
        vals = [
            p.get("code", ""),
            p.get("name", ""),
            p.get("sponsor", ""),
            p.get("pm", ""),
            p.get("phase_or_status", ""),
            p.get("budget", ""),
            p.get("priority", ""),
            p.get("risk_flag", ""),
            p.get("notes", ""),
        ]
        for c, v in enumerate(vals, 1):
            style_xlsx_cell(ws.cell(r, c, v))
    wb.save(dest)


def fill_phase_exit(src: Path, dest: Path, data: dict, meta: dict[str, str]) -> None:
    shutil.copy2(src, dest)
    doc = Document(str(dest))
    replace_in_docx(doc, meta)
    items = (data.get("documents") or {}).get("phase_exit_checklist", {}).get("items") or []
    # first table = checklist
    if doc.tables:
        table = doc.tables[0]
        # clear data rows except header
        while len(table.rows) > 1:
            tbl = table._tbl
            tbl.remove(table.rows[-1]._tr)
        for it in items:
            row = table.add_row().cells
            vals = [
                it.get("item", ""),
                it.get("status", ""),
                it.get("evidence", ""),
                it.get("owner", ""),
            ]
            for cell, v in zip(row, vals):
                cell.paragraphs[0].clear()
                set_para_text(cell.paragraphs[0], v)
    doc.save(dest)


def fill_peer(src: Path, dest: Path, data: dict, meta: dict[str, str]) -> None:
    shutil.copy2(src, dest)
    doc = Document(str(dest))
    replace_in_docx(doc, meta)
    members = (data.get("documents") or {}).get("peer_evaluation", {}).get("members") or []
    if doc.tables:
        table = doc.tables[0]
        while len(table.rows) > 1:
            table._tbl.remove(table.rows[-1]._tr)
        for m in members:
            row = table.add_row().cells
            vals = [
                m.get("name", ""),
                m.get("role", ""),
                m.get("contribution", ""),
                m.get("collaboration", ""),
                m.get("quality", ""),
                m.get("reliability", ""),
                m.get("comment", ""),
            ]
            for cell, v in zip(row, vals):
                cell.paragraphs[0].clear()
                set_para_text(cell.paragraphs[0], "" if v is None else str(v))
    doc.save(dest)


def fill_post_eval(src: Path, dest: Path, data: dict, meta: dict[str, str]) -> None:
    shutil.copy2(src, dest)
    doc = Document(str(dest))
    replace_in_docx(doc, meta)
    doc.save(dest)


def fill_survey(src: Path, dest: Path, data: dict, meta: dict[str, str]) -> None:
    shutil.copy2(src, dest)
    doc = Document(str(dest))
    replace_in_docx(doc, meta)
    answers = (data.get("documents") or {}).get("post_impl_survey", {}).get("answers") or []
    if doc.tables and answers:
        table = doc.tables[0]
        # header + N questions; fill score col
        for i, ans in enumerate(answers):
            ridx = i + 1
            if ridx >= len(table.rows):
                break
            q = ans.get("question")
            score = ans.get("score", "")
            if q:
                table.rows[ridx].cells[0].paragraphs[0].clear()
                set_para_text(table.rows[ridx].cells[0].paragraphs[0], q)
            table.rows[ridx].cells[1].paragraphs[0].clear()
            set_para_text(table.rows[ridx].cells[1].paragraphs[0], "" if score is None else str(score))
    doc.save(dest)


FILLERS = {
    "stakeholder_analysis": fill_stakeholder,
    "issues_actions": fill_issues,
    "selection_matrix": fill_selection,
    "portfolio": fill_portfolio,
    "phase_exit_checklist": fill_phase_exit,
    "peer_evaluation": fill_peer,
    "stakeholder_post_eval": fill_post_eval,
    "post_impl_survey": fill_survey,
}


def fill_one(doc_id: str, data: dict, meta: dict[str, str]) -> Path:
    if doc_id == "project_review_closure":
        raise SystemExit("Dùng: python3 ictpm-flow/closeout/tools/fill_closeout.py")
    if doc_id not in CATALOG:
        raise SystemExit(f"Không biết doc={doc_id!r}. Chạy --list.")
    doc_block = (data.get("documents") or {}).get(doc_id) or {}
    if doc_block.get("enabled") is False:
        print(f"Bỏ qua {doc_id} (enabled=false)")
        return Path()
    tmpl_name, out_name = CATALOG[doc_id]
    src = TEMPLATE_DIR / tmpl_name
    if not src.exists():
        raise SystemExit(f"Thiếu template {src}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    dest = OUT_DIR / out_name
    FILLERS[doc_id](src, dest, data, meta)
    print(f"Đã ghi {dest}")
    return dest


def main() -> int:
    parser = argparse.ArgumentParser(description="Fill supporting toolkit từ 01_supporting.json")
    parser.add_argument("--list", action="store_true")
    parser.add_argument("--doc", help="ID tài liệu")
    parser.add_argument("--all", action="store_true")
    parser.add_argument("--force", action="store_true", help="Fill dù chưa approved")
    parser.add_argument("--json", type=Path, help="Path JSON (mặc định artifacts/01_supporting.json)")
    args = parser.parse_args()

    if args.list or (not args.doc and not args.all):
        print("Catalog:")
        for k, (tmpl, out) in CATALOG.items():
            print(f"  {k:28s} → templates/{tmpl}  ⇒  artifacts/{out}")
        print("  project_review_closure      → pm:closeout")
        if not args.doc and not args.all:
            return 0

    global JSON_PATH
    if args.json:
        JSON_PATH = args.json
    data = load_data()
    if data.get("status") != "approved" and not args.force:
        raise SystemExit(
            f"status={data.get('status')!r} — chỉ fill khi approved (hoặc --force)."
        )
    meta = meta_from(data)
    if args.all:
        requested = (data.get("inputs") or {}).get("docs_requested") or ["all"]
        ids = list(CATALOG.keys()) if requested == ["all"] or requested == "all" else requested
        for doc_id in ids:
            if doc_id in CATALOG:
                fill_one(doc_id, data, meta)
        return 0
    fill_one(args.doc, data, meta)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
