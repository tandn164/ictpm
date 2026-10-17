#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Tạo / làm dày template supporting (TNR 14). Chạy một lần khi nâng cấp."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

MODULE = Path(__file__).resolve().parents[1]
OUT = MODULE / "templates"

FONT = Font(name="Times New Roman", size=14)
FONT_B = Font(name="Times New Roman", size=14, bold=True)
FONT_H = Font(name="Times New Roman", size=16, bold=True)
HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _style_cell(cell, *, bold=False, header=False):
    cell.font = FONT_B if (bold or header) else FONT
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = THIN
    if header:
        cell.fill = HEADER_FILL


def _docx_font(run, size=14, bold=False):
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _add_para(doc, text, *, bold=False, size=14, center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _docx_font(run, size=size, bold=bold)
    return p


def build_stakeholder():
    wb = Workbook()
    ws = wb.active
    ws.title = "Stakeholder Register"
    ws["A1"] = "PHÂN TÍCH STAKEHOLDER"
    ws["A1"].font = FONT_H
    ws.merge_cells("A1:H1")
    ws["A2"] = "Dự án: {{project_name}}  |  Mã: {{project_code}}  |  PM: {{pm}}  |  Ngày: {{date}}"
    ws["A2"].font = FONT
    ws.merge_cells("A2:H2")
    headers = [
        "STT",
        "Họ tên / Đơn vị",
        "Vai trò",
        "Quyền lực",
        "Quan tâm",
        "Chiến lược gắn kết",
        "Owner liên lạc",
        "Ghi chú / Engagement",
    ]
    for i, h in enumerate(headers, 1):
        c = ws.cell(4, i, h)
        _style_cell(c, header=True)
    for r in range(5, 20):
        for c in range(1, 9):
            _style_cell(ws.cell(r, c, ""))
        ws.cell(r, 1, r - 4)
    widths = [6, 28, 28, 12, 12, 22, 16, 32]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    ws2 = wb.create_sheet("Ma trận Quyền lực-Quan tâm")
    ws2["A1"] = "MA TRẬN QUYỀN LỰ – QUAN TÂM (Mendelow)"
    ws2["A1"].font = FONT_H
    ws2.merge_cells("A1:D1")
    ws2["A2"] = "Dự án: {{project_name}}"
    ws2["A2"].font = FONT
    labels = [
        (4, 1, "Quan tâm → / Quyền lực ↓"),
        (4, 2, "Low Interest"),
        (4, 3, "Middle Interest"),
        (4, 4, "High Interest"),
        (5, 1, "High Power"),
        (6, 1, "Middle Power"),
        (7, 1, "Low Power"),
    ]
    for r, c, t in labels:
        cell = ws2.cell(r, c, t)
        _style_cell(cell, header=(r == 4 or c == 1))
    for r in range(5, 8):
        for c in range(2, 5):
            _style_cell(ws2.cell(r, c, ""))
    for i, w in enumerate([28, 28, 28, 28], 1):
        ws2.column_dimensions[chr(64 + i)].width = w
    ws2["A9"] = (
        "Chiến lược gợi ý: High×High = Manage Closely; High Power×Low Interest = Keep Satisfied; "
        "Low Power×High Interest = Keep Informed; Low×Low = Monitor."
    )
    ws2["A9"].font = FONT
    ws2.merge_cells("A9:D9")
    wb.save(OUT / "01_stakeholder_analysis.xlsx")


def build_issues():
    wb = Workbook()
    ws = wb.active
    ws.title = "Issues_Actions"
    ws["A1"] = "CƠ SỞ DỮ LIỆU VẤN ĐỀ & HÀNH ĐỘNG"
    ws["A1"].font = FONT_H
    ws.merge_cells("A1:I1")
    ws["A2"] = "Dự án: {{project_name}}  |  Mã: {{project_code}}  |  PM: {{pm}}  |  Ngày: {{date}}"
    ws["A2"].font = FONT
    ws.merge_cells("A2:I2")
    headers = [
        "ID",
        "Loại",
        "Mô tả",
        "Ưu tiên",
        "Trạng thái",
        "Owner",
        "Hạn",
        "Hành động",
        "Liên quan Risk",
    ]
    for i, h in enumerate(headers, 1):
        _style_cell(ws.cell(4, i, h), header=True)
    for r in range(5, 25):
        for c in range(1, 10):
            _style_cell(ws.cell(r, c, ""))
    for i, w in enumerate([10, 10, 40, 12, 12, 16, 12, 32, 14], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    wb.save(OUT / "02_issues_actions.xlsx")


def build_selection():
    wb = Workbook()
    ws = wb.active
    ws.title = "Selection Matrix"
    ws["A1"] = "MA TRẬN QUYẾT ĐỊNH CHỌN DỰ ÁN"
    ws["A1"].font = FONT_H
    ws.merge_cells("A1:F1")
    ws["A2"] = "Tổ chức: {{company}}  |  Ngày: {{date}}  |  Khuyến nghị: {{selection_recommendation}}"
    ws["A2"].font = FONT
    ws.merge_cells("A2:F2")
    headers = ["Tiêu chí", "Trọng số (0–1)", "Phương án A", "Phương án B", "Phương án C", "Ghi chú"]
    for i, h in enumerate(headers, 1):
        _style_cell(ws.cell(4, i, h), header=True)
    for r in range(5, 15):
        for c in range(1, 7):
            _style_cell(ws.cell(r, c, ""))
    ws["A16"] = "Tổng điểm có trọng số"
    ws["A16"].font = FONT_B
    for c in range(1, 7):
        _style_cell(ws.cell(16, c, ws.cell(16, c).value or ""))
    for i, w in enumerate([28, 14, 14, 14, 14, 28], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    wb.save(OUT / "05_selection_matrix.xlsx")


def build_portfolio():
    wb = Workbook()
    ws = wb.active
    ws.title = "Portfolio"
    ws["A1"] = "QUẢN LÝ DANH MỤC DỰ ÁN (PORTFOLIO)"
    ws["A1"].font = FONT_H
    ws.merge_cells("A1:I1")
    ws["A2"] = "Đơn vị: {{company}}  |  Ngày: {{date}}"
    ws["A2"].font = FONT
    ws.merge_cells("A2:I2")
    headers = [
        "Mã DA",
        "Tên dự án",
        "Sponsor",
        "PM",
        "Giai đoạn / Trạng thái",
        "Ngân sách",
        "Ưu tiên",
        "Cờ rủi ro",
        "Ghi chú",
    ]
    for i, h in enumerate(headers, 1):
        _style_cell(ws.cell(4, i, h), header=True)
    for r in range(5, 20):
        for c in range(1, 10):
            _style_cell(ws.cell(r, c, ""))
    for i, w in enumerate([16, 32, 22, 16, 18, 16, 10, 12, 28], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    wb.save(OUT / "06_portfolio.xlsx")


def build_phase_exit():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2)
        s.right_margin = Cm(2)
    _add_para(doc, "CHECKLIST RA CỔNG GIAI ĐOẠN (Phase Exit Review)", bold=True, size=16, center=True)
    _add_para(
        doc,
        "Dự án: {{project_name}}  |  Mã: {{project_code}}  |  Giai đoạn: {{phase}}  |  Ngày: {{date}}",
    )
    _add_para(doc, "PM: {{pm}}  |  Sponsor: {{sponsor}}")
    _add_para(doc, "1. Mục đích", bold=True)
    _add_para(
        doc,
        "Xác nhận đủ điều kiện chuyển giai đoạn: deliverable, chất lượng, rủi ro/issue mở, "
        "nguồn lực và quyết định cổng (Go / Conditional Go / No-Go).",
    )
    _add_para(doc, "2. Checklist", bold=True)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = ["Hạng mục", "Trạng thái (Done/Open/N/A)", "Bằng chứng", "Owner"]
    for i, h in enumerate(hdr):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        _docx_font(run, bold=True)
    for _ in range(12):
        row = table.add_row().cells
        for c in row:
            c.paragraphs[0].clear()
            run = c.paragraphs[0].add_run("")
            _docx_font(run)
    _add_para(doc, "3. Quyết định cổng", bold=True)
    _add_para(doc, "Quyết định: {{gate_decision}}")
    _add_para(doc, "Điều kiện / hành động còn lại: {{gate_conditions}}")
    _add_para(doc, "Chữ ký PM: __________    Sponsor: __________    Ngày: __________")
    doc.save(OUT / "03_phase_exit_checklist.docx")


def build_peer():
    doc = Document()
    _add_para(doc, "PHIẾU ĐÁNH GIÁ ĐỒNG NGHIỆP TRONG TEAM", bold=True, size=16, center=True)
    _add_para(doc, "Dự án: {{project_name}}  |  Người đánh giá: {{evaluator}}  |  Ngày: {{date}}")
    _add_para(doc, "Thang điểm 1–5 (1=Yếu … 5=Xuất sắc). Để trống nếu chưa có Decision Manual.")
    _add_para(doc, "1. Đánh giá từng thành viên", bold=True)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    hdr = ["Họ tên", "Vai trò", "Đóng góp", "Hợp tác", "Chất lượng", "Tin cậy", "Nhận xét"]
    for i, h in enumerate(hdr):
        run = table.rows[0].cells[i].paragraphs[0].add_run(h)
        _docx_font(run, bold=True)
    for _ in range(10):
        row = table.add_row().cells
        for c in row:
            c.paragraphs[0].clear()
            _docx_font(c.paragraphs[0].add_run(""))
    _add_para(doc, "2. Nhận xét team", bold=True)
    _add_para(doc, "Điểm mạnh: {{team_strengths}}")
    _add_para(doc, "Cần cải thiện: {{team_improve}}")
    _add_para(doc, "Bài học cá nhân: {{personal_lesson}}")
    doc.save(OUT / "04_peer_evaluation.docx")


def build_post_eval():
    doc = Document()
    _add_para(doc, "PHIẾU ĐÁNH GIÁ SAU DỰ ÁN — STAKEHOLDER", bold=True, size=16, center=True)
    _add_para(
        doc,
        "Dự án: {{project_name}}  |  Stakeholder: {{stakeholder_name}}  |  Ngày: {{date}}",
    )
    _add_para(doc, "1. Những gì làm tốt", bold=True)
    _add_para(doc, "{{what_went_well}}")
    _add_para(doc, "2. Thách thức chính", bold=True)
    _add_para(doc, "{{main_challenges}}")
    _add_para(doc, "3. Khuyến nghị", bold=True)
    _add_para(doc, "{{recommendations}}")
    _add_para(doc, "4. Đánh giá nhanh (1–5, để trống nếu chưa có Decision)", bold=True)
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Mức độ đạt mục tiêu dự án", "{{rating_objectives}}"),
        ("Chất lượng giao tiếp", "{{rating_communication}}"),
        ("Quản lý thay đổi / rủi ro", "{{rating_change_risk}}"),
        ("Sẵn sàng vận hành sau bàn giao", "{{rating_ops}}"),
    ]
    for i, (a, b) in enumerate(rows):
        table.rows[i].cells[0].paragraphs[0].clear()
        r0 = table.rows[i].cells[0].paragraphs[0].add_run(a)
        _docx_font(r0, bold=True)
        table.rows[i].cells[1].paragraphs[0].clear()
        r1 = table.rows[i].cells[1].paragraphs[0].add_run(b)
        _docx_font(r1)
    doc.save(OUT / "07_stakeholder_post_evaluation.docx")


def build_survey():
    doc = Document()
    _add_para(doc, "KHẢO SÁT SAU TRIỂN KHAI (Post-Implementation Survey)", bold=True, size=16, center=True)
    _add_para(
        doc,
        "Dự án: {{project_name}}  |  Người trả lời: {{respondent}}  |  Vai trò: {{role}}  |  Ngày: {{date}}",
    )
    _add_para(doc, "Thang Likert 1–5 (1=Hoàn toàn không đồng ý … 5=Hoàn toàn đồng ý). Không bịa điểm.")
    _add_para(doc, "1. Câu hỏi", bold=True)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for i, h in enumerate(["Câu hỏi", "Điểm (1–5)"]):
        run = table.rows[0].cells[i].paragraphs[0].add_run(h)
        _docx_font(run, bold=True)
    questions = [
        "Hệ thống đáp ứng nhu cầu công việc hàng ngày",
        "Hiệu năng / ổn định đạt kỳ vọng",
        "Đào tạo và tài liệu hỗ trợ đủ dùng",
        "Hỗ trợ sau go-live kịp thời",
        "Tôi sẵn sàng giới thiệu giải pháp cho đơn vị khác",
    ]
    for q in questions:
        row = table.add_row().cells
        row[0].paragraphs[0].clear()
        _docx_font(row[0].paragraphs[0].add_run(q))
        row[1].paragraphs[0].clear()
        _docx_font(row[1].paragraphs[0].add_run(""))
    _add_para(doc, "2. Cải thiện ngay", bold=True)
    _add_para(doc, "{{improve_now}}")
    _add_para(doc, "3. Best practice nên giữ", bold=True)
    _add_para(doc, "{{best_practice}}")
    doc.save(OUT / "08_post_implementation_survey.docx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "_source").mkdir(exist_ok=True)
    build_stakeholder()
    build_issues()
    build_selection()
    build_portfolio()
    build_phase_exit()
    build_peer()
    build_post_eval()
    build_survey()
    (OUT / "09_project_review_closure.README.txt").write_text(
        "Project Review & Closure Report → dùng pm:closeout / fill_closeout.py\n",
        encoding="utf-8",
    )
    print("Đã tạo templates trong", OUT)


if __name__ == "__main__":
    main()
