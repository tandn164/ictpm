#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fill masterplan.docx từ 01_masterplan.json (approved).

  python3 ictpm-flow/masterplan/tools/fill_masterplan.py
  python3 ictpm-flow/masterplan/tools/fill_masterplan.py --force
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement

MODULE = Path(__file__).resolve().parents[1]
JSON_PATH = MODULE / "artifacts" / "01_masterplan.json"
TEMPLATE_PATH = MODULE / "templates" / "masterplan.docx"
OUT_PATH = MODULE / "artifacts" / "01_masterplan_draft.docx"

PLACEHOLDER_RE = re.compile(r"\{\{([a-z0-9_]+)\}\}")


def as_plain_text(cur, default: str = "") -> str:
    if cur is None:
        return default
    if isinstance(cur, dict):
        if "value" in cur or "proposed" in cur:
            v = cur.get("value")
            if v is None or (isinstance(v, str) and not str(v).strip()):
                v = cur.get("proposed")
            if v is None:
                return default
            if isinstance(v, (list, tuple)):
                return "\n".join(f"• {x}" for x in v)
            if isinstance(v, dict):
                return as_plain_text(v, default=default)
            return str(v).strip()
        return default
    if isinstance(cur, (list, tuple)):
        return "\n".join(f"• {x}" for x in cur)
    return str(cur).strip()


def _rfonts(run, size: int = 14, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")


def set_cell_text(cell, text: str, *, bold: bool = False, size: int = 12) -> None:
    text = "" if text is None else str(text)
    lines = text.split("\n") if text else [""]
    while len(cell.paragraphs) > len(lines):
        p = cell.paragraphs[-1]._p
        p.getparent().remove(p)
    while len(cell.paragraphs) < len(lines):
        cell.add_paragraph()
    for i, line in enumerate(lines):
        p = cell.paragraphs[i]
        for r in p.runs:
            r.text = ""
        if p.runs:
            p.runs[0].text = line
            _rfonts(p.runs[0], size=size, bold=bold)
        else:
            run = p.add_run(line)
            _rfonts(run, size=size, bold=bold)


def shade_header_row(table) -> None:
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        # remove old shd
        for child in list(tcPr):
            if child.tag == qn("w:shd"):
                tcPr.remove(child)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1F4E79")
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                _rfonts(r, size=12, bold=True)


def replace_placeholders_in_paragraph(paragraph, mapping: dict[str, str]) -> None:
    if not paragraph.runs and not paragraph.text:
        return
    full = paragraph.text
    if "{{" not in full:
        return

    def repl(m: re.Match) -> str:
        key = m.group(1)
        return mapping.get(key, m.group(0))

    new_text = PLACEHOLDER_RE.sub(repl, full)
    if new_text == full:
        return
    # rewrite first run, clear others
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        _rfonts(paragraph.runs[0], size=14)
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        run = paragraph.add_run(new_text)
        _rfonts(run, size=14)


def find_marker_table(doc: Document, marker: str):
    """Tìm bảng có ô chứa {{marker}}."""
    token = "{{" + marker + "}}"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if token in cell.text:
                    return table
    return None


def refill_table(table, headers: list[str], rows: list[list[str]]) -> None:
    """Giữ hàng header; xóa hết hàng data; thêm rows mới."""
    # delete data rows (keep row 0)
    while len(table.rows) > 1:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)
    # ensure header text
    for i, h in enumerate(headers):
        if i < len(table.columns):
            set_cell_text(table.rows[0].cells[i], h, bold=True, size=12)
    shade_header_row(table)
    if not rows:
        row = table.add_row()
        set_cell_text(row.cells[0], "—")
        for c in range(1, len(headers)):
            set_cell_text(row.cells[c], "")
        return
    for data in rows:
        row = table.add_row()
        for i, val in enumerate(data):
            if i < len(headers):
                set_cell_text(row.cells[i], val or "", size=12)


def fill_from_json(data: dict, template: Path, out: Path) -> None:
    shutil.copy2(template, out)
    doc = Document(str(out))

    project = data.get("project") or {}
    overview = data.get("overview") or {}

    scalar = {
        "company": as_plain_text(project.get("company")),
        "project_name": as_plain_text(project.get("name")) or as_plain_text(data.get("project_name")),
        "project_code": as_plain_text(project.get("code")),
        "pm": as_plain_text(project.get("pm")),
        "date": as_plain_text(project.get("date")),
        "version": as_plain_text(project.get("version")) or "1.0",
        "project_summary": as_plain_text(overview.get("project_summary")),
        "scope_in": as_plain_text(overview.get("scope_in")),
        "scope_out": as_plain_text(overview.get("scope_out")),
        "assumptions": as_plain_text(overview.get("assumptions")),
        "constraints": as_plain_text(overview.get("constraints")),
        "management_approach": as_plain_text(overview.get("management_approach")),
        "roles_intro": as_plain_text(data.get("roles_intro")),
        "communications_intro": as_plain_text(data.get("communications_intro")),
        "schedule_intro": as_plain_text(data.get("schedule_intro")),
        "issues_intro": as_plain_text(data.get("issues_intro")),
        "risk_summary": as_plain_text(data.get("risk_summary")),
        "transition_intro": as_plain_text(data.get("transition_intro")),
        "cost_summary": as_plain_text(data.get("cost_summary")),
        "quality_summary": as_plain_text(data.get("quality_summary")),
    }

    for p in doc.paragraphs:
        replace_placeholders_in_paragraph(p, scalar)
    # also version history first data row may have placeholders in cells
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_placeholders_in_paragraph(p, scalar)

    # Roles
    roles = data.get("roles") or []
    t = find_marker_table(doc, "roles_table")
    if t is not None:
        refill_table(
            t,
            ["Họ tên", "Vai trò", "Trách nhiệm"],
            [
                [
                    as_plain_text(r.get("name")),
                    as_plain_text(r.get("role")),
                    as_plain_text(r.get("responsibility")),
                ]
                for r in roles
            ],
        )

    # Communications
    comms = data.get("communications") or []
    t = find_marker_table(doc, "communications_table")
    if t is not None:
        refill_table(
            t,
            ["Nhóm stakeholder", "Nội dung giao tiếp", "Mục đích"],
            [
                [
                    as_plain_text(c.get("stakeholder_group")),
                    as_plain_text(c.get("item")),
                    as_plain_text(c.get("purpose")),
                ]
                for c in comms
            ],
        )

    # Schedule
    schedule = data.get("schedule") or []
    t = find_marker_table(doc, "schedule_table")
    if t is not None:
        refill_table(
            t,
            ["Mốc / Công việc", "Bắt đầu (KH)", "Kết thúc (KH)", "% / Thực tế", "Owner"],
            [
                [
                    as_plain_text(s.get("milestone")),
                    as_plain_text(s.get("start")),
                    as_plain_text(s.get("end")),
                    as_plain_text(s.get("actual_or_pct")),
                    as_plain_text(s.get("owner")),
                ]
                for s in schedule
            ],
        )

    # Issues
    issues = data.get("issues") or []
    t = find_marker_table(doc, "issues_table")
    if t is not None:
        refill_table(
            t,
            ["Tiêu đề", "Mô tả", "Ngày phát hiện", "Người nêu", "Mức ảnh hưởng", "Hướng xử lý"],
            [
                [
                    as_plain_text(i.get("title")),
                    as_plain_text(i.get("description")),
                    as_plain_text(i.get("date_identified")),
                    as_plain_text(i.get("originator")),
                    as_plain_text(i.get("impact")),
                    as_plain_text(i.get("resolution")),
                ]
                for i in issues
            ]
            if issues
            else [["(Chưa có issue)", "", "", "", "", ""]],
        )

    # Risks top
    risks = data.get("risks_top") or []
    t = find_marker_table(doc, "risks_table")
    if t is not None:
        refill_table(
            t,
            ["ID / Tiêu đề", "Mô tả", "Ngày", "Originator", "Priority (A/B/C)", "Ứng phó"],
            [
                [
                    (as_plain_text(r.get("id")) + " " + as_plain_text(r.get("title"))).strip(),
                    as_plain_text(r.get("description")) or as_plain_text(r.get("risk")),
                    as_plain_text(r.get("date_identified")),
                    as_plain_text(r.get("originator")),
                    as_plain_text(r.get("priority")),
                    as_plain_text(r.get("response"))
                    or (
                        as_plain_text(r.get("strategy"))
                        + (
                            ": " + as_plain_text(r.get("response_actions"))
                            if as_plain_text(r.get("response_actions"))
                            else ""
                        )
                    ).strip(": "),
                ]
                for r in risks
            ],
        )

    # Transition
    transition = data.get("transition") or []
    t = find_marker_table(doc, "transition_table")
    if t is not None:
        refill_table(
            t,
            ["Hạng mục sẵn sàng", "Có/Không", "Biện pháp nếu Không"],
            [
                [
                    as_plain_text(x.get("item")),
                    as_plain_text(x.get("ready")),
                    as_plain_text(x.get("mitigation_if_no")),
                ]
                for x in transition
            ],
        )

    # Cost
    cost_items = data.get("cost_items") or []
    t = find_marker_table(doc, "cost_table")
    if t is not None:
        refill_table(
            t,
            ["Hạng mục", "Số tiền / tỷ lệ", "Ghi chú"],
            [
                [
                    as_plain_text(x.get("item")),
                    as_plain_text(x.get("amount")),
                    as_plain_text(x.get("note")),
                ]
                for x in cost_items
            ]
            if cost_items
            else [["(Chờ Decision ngân sách)", "", ""]],
        )

    # Quality
    quality = data.get("quality_criteria") or []
    t = find_marker_table(doc, "quality_table")
    if t is not None:
        refill_table(
            t,
            ["Tiêu chí / Deliverable", "Cách kiểm tra", "Điều kiện chấp nhận"],
            [
                [
                    as_plain_text(x.get("deliverable")),
                    as_plain_text(x.get("how_tested")),
                    as_plain_text(x.get("acceptance")),
                ]
                for x in quality
            ],
        )

    # leftover placeholders → empty
    for p in doc.paragraphs:
        if "{{" in p.text:
            replace_placeholders_in_paragraph(p, {k: "" for k in PLACEHOLDER_RE.findall(p.text)})
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if "{{" in p.text:
                        replace_placeholders_in_paragraph(
                            p, {k: "" for k in PLACEHOLDER_RE.findall(p.text)}
                        )

    doc.save(out)
    print(f"Đã ghi {out}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Fill Master Plan Word từ JSON")
    parser.add_argument("--force", action="store_true", help="Fill dù chưa approved")
    args = parser.parse_args()

    if not TEMPLATE_PATH.exists():
        print(f"Thiếu template {TEMPLATE_PATH}", file=sys.stderr)
        return 1
    if not JSON_PATH.exists():
        print(f"Thiếu {JSON_PATH}", file=sys.stderr)
        return 1

    data = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    if data.get("status") != "approved" and not args.force:
        print(
            f"Từ chối fill: status={data.get('status')!r}. "
            "Chỉ chạy khi status=approved (hoặc --force).",
            file=sys.stderr,
        )
        return 2

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    fill_from_json(data, TEMPLATE_PATH, OUT_PATH)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
