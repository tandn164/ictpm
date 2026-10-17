#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fill closeout.docx từ 01_closeout.json (approved).

  python3 ictpm-flow/closeout/tools/fill_closeout.py
  python3 ictpm-flow/closeout/tools/fill_closeout.py --force
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

MODULE = Path(__file__).resolve().parents[1]
JSON_PATH = MODULE / "artifacts" / "01_closeout.json"
TEMPLATE_PATH = MODULE / "templates" / "closeout.docx"
OUT_PATH = MODULE / "artifacts" / "01_closeout_draft.docx"

PLACEHOLDER_RE = re.compile(r"\{\{([a-z0-9_]+)\}\}")


def as_plain_text(cur, default: str = "") -> str:
    if cur is None:
        return default
    if isinstance(cur, dict):
        if "value" in cur or "proposed" in cur:
            v = cur.get("value")
            if v is None or (isinstance(v, str) and not str(v).strip()):
                v = cur.get("proposed")
            if v is None:
                return default
            if isinstance(v, (list, tuple)):
                return "\n".join(f"• {x}" for x in v)
            if isinstance(v, dict):
                return as_plain_text(v, default=default)
            return str(v).strip()
        return default
    if isinstance(cur, (list, tuple)):
        return "\n".join(f"• {x}" for x in cur)
    return str(cur).strip()


def _rfonts(run, size: int = 14, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")


def set_cell_text(cell, text: str, *, bold: bool = False, size: int = 12) -> None:
    text = "" if text is None else str(text)
    lines = text.split("\n") if text else [""]
    while len(cell.paragraphs) > len(lines):
        p = cell.paragraphs[-1]._p
        p.getparent().remove(p)
    while len(cell.paragraphs) < len(lines):
        cell.add_paragraph()
    for i, line in enumerate(lines):
        p = cell.paragraphs[i]
        for r in p.runs:
            r.text = ""
        if p.runs:
            p.runs[0].text = line
            _rfonts(p.runs[0], size=size, bold=bold)
        else:
            run = p.add_run(line)
            _rfonts(run, size=size, bold=bold)


def shade_header_row(table) -> None:
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        for child in list(tcPr):
            if child.tag == qn("w:shd"):
                tcPr.remove(child)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1F4E79")
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                _rfonts(r, size=12, bold=True)


def replace_placeholders_in_paragraph(paragraph, mapping: dict[str, str]) -> None:
    full = paragraph.text
    if "{{" not in full:
        return

    def repl(m: re.Match) -> str:
        return mapping.get(m.group(1), m.group(0))

    new_text = PLACEHOLDER_RE.sub(repl, full)
    if new_text == full:
        return
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        _rfonts(paragraph.runs[0], size=14)
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        run = paragraph.add_run(new_text)
        _rfonts(run, size=14)


def find_marker_table(doc: Document, marker: str):
    token = "{{" + marker + "}}"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if token in cell.text:
                    return table
    return None


def refill_table(table, headers: list[str], rows: list[list[str]]) -> None:
    while len(table.rows) > 1:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)
    for i, h in enumerate(headers):
        if i < len(table.columns):
            set_cell_text(table.rows[0].cells[i], h, bold=True, size=12)
    shade_header_row(table)
    if not rows:
        row = table.add_row()
        set_cell_text(row.cells[0], "—")
        for c in range(1, len(headers)):
            set_cell_text(row.cells[c], "")
        return
    for data in rows:
        row = table.add_row()
        for i, val in enumerate(data):
            if i < len(headers):
                set_cell_text(row.cells[i], val or "", size=12)


def fill_from_json(data: dict, template: Path, out: Path) -> None:
    shutil.copy2(template, out)
    doc = Document(str(out))

    project = data.get("project") or {}
    eff = data.get("effectiveness") or {}
    admin = data.get("admin_closure") or {}
    contract = data.get("contract_closure") or {}
    archive = data.get("archive") or {}

    scalar = {
        "company": as_plain_text(project.get("company")),
        "project_name": as_plain_text(project.get("name")) or as_plain_text(data.get("project_name")),
        "project_code": as_plain_text(project.get("code")),
        "pm": as_plain_text(project.get("pm")),
        "sponsor": as_plain_text(project.get("sponsor")),
        "date": as_plain_text(project.get("date")),
        "version": as_plain_text(project.get("version")) or "1.0",
        "executive_summary": as_plain_text(data.get("executive_summary")),
        "objectives_met": as_plain_text(data.get("objectives_met")),
        "general_comments": as_plain_text(data.get("general_comments")),
        "effectiveness_summary": as_plain_text(eff.get("summary")),
        "lessons_intro": as_plain_text(data.get("lessons_intro")),
        "admin_summary": as_plain_text(admin.get("summary")),
        "contract_summary": as_plain_text(contract.get("summary")),
        "archive_summary": as_plain_text(archive.get("summary")),
        "references": as_plain_text(data.get("references")),
    }

    for p in doc.paragraphs:
        replace_placeholders_in_paragraph(p, scalar)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_placeholders_in_paragraph(p, scalar)

    baselines = eff.get("baselines") or []
    t = find_marker_table(doc, "baselines_table")
    if t is not None:
        refill_table(
            t,
            ["Lĩnh vực", "Kế hoạch (baseline)", "Thực tế", "Chênh lệch / ghi chú"],
            [
                [
                    as_plain_text(b.get("area")),
                    as_plain_text(b.get("planned")),
                    as_plain_text(b.get("actual")) or "(chờ Fact/Decision)",
                    as_plain_text(b.get("variance")),
                ]
                for b in baselines
            ],
        )

    lessons = data.get("lessons") or []
    t = find_marker_table(doc, "lessons_table")
    if t is not None:
        refill_table(
            t,
            ["Nhóm", "Sự việc", "Nguyên nhân", "Khuyến nghị", "Mức"],
            [
                [
                    as_plain_text(x.get("category")),
                    as_plain_text(x.get("what_happened")),
                    as_plain_text(x.get("root_cause")),
                    as_plain_text(x.get("recommendation")),
                    as_plain_text(x.get("priority")),
                ]
                for x in lessons
            ]
            if lessons
            else [["—", "(Chưa có lesson — chờ Manual)", "", "", ""]],
        )

    deliverables = admin.get("deliverables") or []
    t = find_marker_table(doc, "deliverables_table")
    if t is not None:
        refill_table(
            t,
            ["Deliverable", "Chấp nhận", "Minh chứng", "Ngày"],
            [
                [
                    as_plain_text(x.get("name")),
                    as_plain_text(x.get("accepted")),
                    as_plain_text(x.get("evidence")),
                    as_plain_text(x.get("date")),
                ]
                for x in deliverables
            ]
            if deliverables
            else [["(Chờ nghiệm thu)", "", "", ""]],
        )

    open_items = admin.get("open_items") or []
    t = find_marker_table(doc, "open_items_table")
    if t is not None:
        refill_table(
            t,
            ["Hạng mục", "Owner", "Hạn", "Trạng thái"],
            [
                [
                    as_plain_text(x.get("item")),
                    as_plain_text(x.get("owner")),
                    as_plain_text(x.get("due")),
                    as_plain_text(x.get("status")),
                ]
                for x in open_items
            ]
            if open_items
            else [["(Không có / chưa liệt kê)", "", "", ""]],
        )

    contracts = contract.get("contracts") or []
    t = find_marker_table(doc, "contracts_table")
    if t is not None:
        refill_table(
            t,
            ["Hợp đồng / gói", "Trạng thái", "Ghi chú"],
            [
                [
                    as_plain_text(x.get("name")),
                    as_plain_text(x.get("status")),
                    as_plain_text(x.get("note")),
                ]
                for x in contracts
            ]
            if contracts
            else [["(Chờ Decision đóng HĐ)", "", ""]],
        )

    items = archive.get("items") or []
    t = find_marker_table(doc, "archive_table")
    if t is not None:
        refill_table(
            t,
            ["Tài liệu / artifact", "Vị trí lưu", "Phân phối tới"],
            [
                [
                    as_plain_text(x.get("item")),
                    as_plain_text(x.get("location")),
                    as_plain_text(x.get("distributed_to")),
                ]
                for x in items
            ]
            if items
            else [["(Đề xuất lưu ICTPM artifacts)", "ictpm-flow/*/artifacts/", "PM / Sponsor"]],
        )

    for p in doc.paragraphs:
        if "{{" in p.text:
            replace_placeholders_in_paragraph(p, {k: "" for k in PLACEHOLDER_RE.findall(p.text)})
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if "{{" in p.text:
                        replace_placeholders_in_paragraph(
                            p, {k: "" for k in PLACEHOLDER_RE.findall(p.text)}
                        )

    doc.save(out)
    print(f"Đã ghi {out}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Fill Báo cáo tổng kết dự án từ JSON")
    parser.add_argument("--force", action="store_true", help="Fill dù chưa approved")
    args = parser.parse_args()

    if not TEMPLATE_PATH.exists():
        print(f"Thiếu template {TEMPLATE_PATH}", file=sys.stderr)
        return 1
    if not JSON_PATH.exists():
        print(f"Thiếu {JSON_PATH}", file=sys.stderr)
        return 1

    data = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    if data.get("status") != "approved" and not args.force:
        print(
            f"Từ chối fill: status={data.get('status')!r}. "
            "Chỉ chạy khi status=approved (hoặc --force).",
            file=sys.stderr,
        )
        return 2

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    fill_from_json(data, TEMPLATE_PATH, OUT_PATH)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
