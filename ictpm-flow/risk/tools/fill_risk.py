#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fill risk-management.xlsx từ 01_risk.json (approved).

Cũng có thể tạo lại template:
  python3 ictpm-flow/risk/tools/fill_risk.py --build-template
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MODULE = Path(__file__).resolve().parents[1]
JSON_PATH = MODULE / "artifacts" / "01_risk.json"
TEMPLATE_PATH = MODULE / "templates" / "risk-management.xlsx"
OUT_PATH = MODULE / "artifacts" / "01_risk_draft.xlsx"

NAVY = "1F4E79"
NAVY2 = "2E75B6"
GOLD = "FFEB9C"
GOLD_FONT = "9C5700"
GREEN = "C6EFCE"
GREEN_FONT = "006100"
GRAY = "F2F2F2"
LIGHT = "D6EAF8"
RED = "FFC7CE"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

PRIORITY_MATRIX = {
    ("High", "Big"): "A",
    ("High", "Middle"): "A",
    ("High", "Small"): "B",
    ("Middle", "Big"): "A",
    ("Middle", "Middle"): "B",
    ("Middle", "Small"): "C",
    ("Low", "Big"): "B",
    ("Low", "Middle"): "C",
    ("Low", "Small"): "C",
}


def fill_color(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def font(bold: bool = False, size: int = 14, color: str = "000000", name: str = "Times New Roman") -> Font:
    return Font(name=name, bold=bold, size=size, color=color)


def clear_fill(cell) -> None:
    """Ô dữ liệu / cột thành phần: không tô màu."""
    cell.fill = PatternFill(fill_type=None)


def align(h: str = "left", v: str = "center", wrap: bool = True) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def as_plain_text(cur, default: str = "") -> str:
    if cur is None:
        return default
    if isinstance(cur, dict):
        if "value" in cur or "proposed" in cur:
            v = cur.get("value")
            if v is None or (isinstance(v, str) and not str(v).strip()):
                v = cur.get("proposed")
            if v is None:
                return default
            if isinstance(v, (list, tuple)):
                return ", ".join(str(x) for x in v)
            if isinstance(v, dict):
                return as_plain_text(v, default=default)
            return str(v).strip()
        return default
    if isinstance(cur, (list, tuple)):
        return ", ".join(str(x) for x in cur)
    return str(cur).strip()


def compute_priority(probability: str, impact: str) -> str:
    return PRIORITY_MATRIX.get((probability.strip(), impact.strip()), "")


def set_cell(ws, row: int, col: int, value, *, bold=False, fill=None, center=False):
    cell = ws.cell(row, col, value)
    cell.font = font(bold=bold, size=14)
    cell.border = THIN
    cell.alignment = align("center" if center else "left", "center", True)
    if fill:
        cell.fill = fill_color(fill)
    else:
        clear_fill(cell)
    return cell


# ---------------------------------------------------------------------------
# Template builder
# ---------------------------------------------------------------------------
def build_template(path: Path = TEMPLATE_PATH) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # --- Cover ---
    ws = wb.active
    ws.title = "Cover"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([3, 22, 28, 18, 18, 14, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.merge_cells("B2:F2")
    ws["B2"] = "PROJECT RISK MANAGEMENT"
    ws["B2"].font = font(True, 20, NAVY)
    ws.merge_cells("B3:F3")
    ws["B3"] = "Plan · Identify · Analyze (A/B/C) · Respond · Monitor"
    ws["B3"].font = font(False, 14, NAVY2)

    ws["B5"] = "Thông tin dự án"
    ws["B5"].font = font(True, 14, WHITE)
    ws["B5"].fill = fill_color(NAVY)
    for c in range(2, 7):
        ws.cell(5, c).fill = fill_color(NAVY)

    meta = [
        (6, "Tên dự án", "project_name"),
        (7, "Mã / Document No.", "project_code"),
        (8, "Công ty / Đơn vị", "company"),
        (9, "Project Manager", "pm"),
        (10, "Ngày lập", date.today().isoformat()),
        (11, "Phiên bản", "1.0"),
    ]
    for r, lab, val in meta:
        ws.cell(r, 2, lab).font = font(True, 14, NAVY)
        ws.cell(r, 2).fill = fill_color(GRAY)
        ws.cell(r, 2).border = THIN
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        cell = ws.cell(r, 3, val if r >= 10 else "")
        cell.font = font(False, 14)
        for c in range(3, 6):
            clear_fill(ws.cell(r, c))
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(False, 14)

    ws["B13"] = "Sheets"
    ws["B13"].font = font(True, 14, WHITE)
    for c in range(2, 5):
        ws.cell(13, c).fill = fill_color(NAVY2)
    sheets_help = [
        ("Plan", "Methodology, Roles, Budgeting, Timing, Categories"),
        ("Register", "Sổ rủi ro + ma trận Priority A/B/C (cùng sheet)"),
        ("Values", "Dropdown Probability / Impact / Strategy / Category"),
        ("Analysis", "Mẫu phiếu — fill tạo thêm 1 sheet / rủi ro (R-001, R-002, …)"),
    ]
    for i, (a, b) in enumerate(sheets_help, 14):
        cell_a = ws.cell(i, 2, a)
        cell_a.border = THIN
        cell_a.font = font(True, 14)
        clear_fill(cell_a)
        ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=5)
        cell_b = ws.cell(i, 3, b)
        cell_b.border = THIN
        cell_b.font = font(False, 14)
        clear_fill(cell_b)

    # --- Plan ---
    ws = wb.create_sheet("Plan")
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([22, 70], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.merge_cells("A1:B1")
    ws["A1"] = "RISK MANAGEMENT PLAN"
    ws["A1"].font = font(True, 16, WHITE)
    ws["A1"].fill = fill_color(NAVY)
    ws["B1"].fill = fill_color(NAVY)
    ws.row_dimensions[1].height = 26

    ws["A2"] = "Project"
    ws["A2"].font = font(True, 14)
    ws["A2"].fill = fill_color(GRAY)
    ws["A2"].border = THIN
    ws["B2"] = "='Cover'!C6"
    clear_fill(ws["B2"])
    ws["B2"].font = font(False, 14)
    ws["B2"].border = THIN

    plan_sections = [
        (4, "1. Methodology", "[Nhập phương pháp QL rủi ro…]"),
        (7, "2. Roles and responsibilities", "[Nhập vai trò…]"),
        (10, "3. Budgeting", "[Nhập ngân sách / dự phòng — chờ Decision nếu thiếu số…]"),
        (13, "4. Timing", "[Nhập lịch nhận diện / rà soát…]"),
        (16, "5. Risk categories", "System, Technical, Organization"),
        (19, "6. Monitor (review questions)", ""),
    ]
    for start, title, placeholder in plan_sections:
        ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=2)
        ws.cell(start, 1, title).font = font(True, 14, WHITE)
        ws.cell(start, 1).fill = fill_color(NAVY2)
        ws.cell(start, 2).fill = fill_color(NAVY2)
        ws.merge_cells(start_row=start + 1, start_column=1, end_row=start + 2, end_column=2)
        cell = ws.cell(start + 1, 1, placeholder)
        clear_fill(cell)
        cell.font = font(False, 14)
        cell.alignment = align("left", "top", True)
        cell.border = THIN
        for r in (start + 1, start + 2):
            for c in (1, 2):
                ws.cell(r, c).border = THIN
                clear_fill(ws.cell(r, c))
                ws.cell(r, c).font = font(False, 14)

    default_q = (
        "Response plan còn hiệu lực?\n"
        "Có rủi ro mới không?\n"
        "Rủi ro đã đổi Probability / Impact / Priority không?\n"
        "Contingency đã sẵn sàng khi trigger?"
    )
    ws.cell(20, 1, default_q)
    ws.cell(20, 1).font = font(False, 14)

    prio_fill = {"A": RED, "B": ORANGE, "C": GREEN}

    # --- Values ---
    ws = wb.create_sheet("Values")
    headers = ["Probability", "Impact", "Strategy", "Category", "Priority"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = font(True, 14, WHITE)
        cell.fill = fill_color(NAVY)
        cell.border = THIN
        ws.column_dimensions[get_column_letter(c)].width = 16
    for i, v in enumerate(["High", "Middle", "Low"], 2):
        cell = ws.cell(i, 1, v)
        cell.border = THIN
        cell.font = font(False, 14)
        clear_fill(cell)
    for i, v in enumerate(["Big", "Middle", "Small"], 2):
        cell = ws.cell(i, 2, v)
        cell.border = THIN
        cell.font = font(False, 14)
        clear_fill(cell)
    for i, v in enumerate(["Avoid", "Transfer", "Mitigate", "Accept"], 2):
        cell = ws.cell(i, 3, v)
        cell.border = THIN
        cell.font = font(False, 14)
        clear_fill(cell)
    for i, v in enumerate(["System", "Technical", "Organization"], 2):
        cell = ws.cell(i, 4, v)
        cell.border = THIN
        cell.font = font(False, 14)
        clear_fill(cell)
    for i, v in enumerate(["A", "B", "C"], 2):
        cell = ws.cell(i, 5, v)
        cell.border = THIN
        cell.font = font(True, 14)
        cell.fill = fill_color(prio_fill[v])
        cell.alignment = align("center", "center", False)

    # --- Register (+ Priority Matrix trên cùng sheet) ---
    ws = wb.create_sheet("Register")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:L1")
    ws["A1"] = "RISK REGISTER"
    ws["A1"].font = font(True, 16, WHITE)
    for c in range(1, 13):
        ws.cell(1, c).fill = fill_color(NAVY)
    ws.row_dimensions[1].height = 24

    ws["A2"] = "Project"
    ws["A2"].font = font(True, 14)
    ws["B2"] = "='Cover'!C6"
    clear_fill(ws["B2"])
    ws["B2"].font = font(False, 14)
    ws["A2"].border = THIN
    ws["B2"].border = THIN

    # Ma trận ưu tiên — phía trên sổ rủi ro để đối chiếu khi chấm P/I
    ws.merge_cells("A4:D4")
    ws["A4"] = "PRIORITY MATRIX (Probability × Impact)"
    ws["A4"].font = font(True, 14, WHITE)
    for c in range(1, 5):
        ws.cell(4, c).fill = fill_color(NAVY2)

    ws["A5"] = "Probability \\ Impact"
    ws["B5"] = "Big"
    ws["C5"] = "Middle"
    ws["D5"] = "Small"
    for c in range(1, 5):
        ws.cell(5, c).font = font(True, 14, WHITE)
        ws.cell(5, c).fill = fill_color(NAVY)
        ws.cell(5, c).border = THIN
        ws.cell(5, c).alignment = align("center", "center", False)

    matrix_rows = [
        ("High", "A", "A", "B"),
        ("Middle", "A", "B", "C"),
        ("Low", "B", "C", "C"),
    ]
    for i, (p, *vals) in enumerate(matrix_rows, 6):
        ws.cell(i, 1, p).font = font(True, 14)
        ws.cell(i, 1).fill = fill_color(GRAY)
        ws.cell(i, 1).border = THIN
        ws.cell(i, 1).alignment = align("center", "center", False)
        for j, v in enumerate(vals, 2):
            cell = ws.cell(i, j, v)
            cell.font = font(True, 14)
            cell.fill = fill_color(prio_fill[v])  # Priority: tô theo mức A/B/C
            cell.border = THIN
            cell.alignment = align("center", "center", False)

    # Chú thích thang — cạnh ma trận
    ws.merge_cells("F4:L4")
    ws["F4"] = "Định nghĩa thang đánh giá"
    ws["F4"].font = font(True, 14, WHITE)
    for c in range(6, 13):
        ws.cell(4, c).fill = fill_color(NAVY2)

    legend_rows = [
        ("P: High / Middle / Low", "Nhiều lần trong quá khứ / vài lần / hiếm hoặc chưa"),
        ("I: Big", "Tổ chức phải xử lý (ảnh hưởng lớn, pháp lý, truyền thông)"),
        ("I: Middle", "Cần team hoặc đơn vị khác hỗ trợ"),
        ("I: Small", "Team dự án tự xử trong buffer"),
        ("A / B / C", "Ưu tiên ứng phó trước / có phương án theo dõi / theo dõi nhẹ"),
    ]
    for i, (a, b) in enumerate(legend_rows, 5):
        ws.cell(i, 6, a).font = font(True, 14)
        ws.cell(i, 6).fill = fill_color(GRAY)
        ws.cell(i, 6).border = THIN
        ws.merge_cells(start_row=i, start_column=7, end_row=i, end_column=12)
        ws.cell(i, 7, b).border = THIN
        ws.cell(i, 7).font = font(False, 14)
        ws.cell(i, 7).fill = fill_color(LIGHT)
        for c in range(7, 13):
            ws.cell(i, c).border = THIN
            ws.cell(i, c).fill = fill_color(LIGHT)
            ws.cell(i, c).font = font(False, 14)

    # Bảng sổ rủi ro — có Initial / Current như RiskRegisterTemplate
    # Row 10: nhóm Initial | Current
    # Row 11: header chi tiết
    # Row 12+: data
    group_row, header_row = 10, 11
    ncols = 15

    def prio_formula(p_col: str, i_col: str, r: int) -> str:
        return (
            f'=IF(OR({p_col}{r}="",{i_col}{r}=""),"",'
            f'IF(AND({p_col}{r}="High",{i_col}{r}="Big"),"A",'
            f'IF(AND({p_col}{r}="High",{i_col}{r}="Middle"),"A",'
            f'IF(AND({p_col}{r}="High",{i_col}{r}="Small"),"B",'
            f'IF(AND({p_col}{r}="Middle",{i_col}{r}="Big"),"A",'
            f'IF(AND({p_col}{r}="Middle",{i_col}{r}="Middle"),"B",'
            f'IF(AND({p_col}{r}="Middle",{i_col}{r}="Small"),"C",'
            f'IF(AND({p_col}{r}="Low",{i_col}{r}="Big"),"B",'
            f'IF(AND({p_col}{r}="Low",{i_col}{r}="Middle"),"C",'
            f'IF(AND({p_col}{r}="Low",{i_col}{r}="Small"),"C",""))))))))))'
        )

    # Group headers (background nhóm Initial/Current — được tô màu)
    for c in range(1, 8):
        ws.cell(group_row, c).fill = fill_color(NAVY2)
        ws.cell(group_row, c).border = THIN
    ws.merge_cells("H10:J10")
    ws["H10"] = "Initial (trước ứng phó)"
    ws["H10"].font = font(True, 14, GOLD_FONT)
    ws["H10"].fill = fill_color(GOLD)
    ws["H10"].alignment = align("center", "center", False)
    for c in range(8, 11):
        ws.cell(group_row, c).fill = fill_color(GOLD)
        ws.cell(group_row, c).border = THIN
    ws.merge_cells("K10:M10")
    ws["K10"] = "Current (sau ứng phó)"
    ws["K10"].font = font(True, 14, GREEN_FONT)
    ws["K10"].fill = fill_color(GREEN)
    ws["K10"].alignment = align("center", "center", False)
    for c in range(11, 14):
        ws.cell(group_row, c).fill = fill_color(GREEN)
        ws.cell(group_row, c).border = THIN
    for c in range(14, 16):
        ws.cell(group_row, c).fill = fill_color(NAVY2)
        ws.cell(group_row, c).border = THIN

    reg_headers = [
        "ID",  # A
        "Category",  # B
        "Risk Driver",  # C
        "Risk",  # D
        "Strategy",  # E
        "Response actions",  # F
        "Owner",  # G
        "P",  # H Initial
        "I",  # I
        "Priority",  # J
        "P",  # K Current
        "I",  # L
        "Priority",  # M
        "Contingency / Residual / Secondary",  # N
        "Evidence / Comments",  # O
    ]
    widths = [8, 12, 20, 26, 11, 24, 10, 8, 8, 10, 8, 8, 10, 26, 20]
    # Header cột thành phần: chỉ nền header (navy), không tô vàng/vàng-xanh từng cột
    for c, (h, w) in enumerate(zip(reg_headers, widths), 1):
        cell = ws.cell(header_row, c, h)
        cell.border = THIN
        cell.alignment = align("center", "center", True)
        cell.font = font(True, 14, WHITE)
        cell.fill = fill_color(NAVY2)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[header_row].height = 28
    ws.freeze_panes = "A12"

    dv_p = DataValidation(type="list", formula1="=Values!$A$2:$A$4", allow_blank=True)
    dv_i = DataValidation(type="list", formula1="=Values!$B$2:$B$4", allow_blank=True)
    dv_s = DataValidation(type="list", formula1="=Values!$C$2:$C$5", allow_blank=True)
    dv_c = DataValidation(type="list", formula1="=Values!$D$2:$D$20", allow_blank=True)
    for dv in (dv_p, dv_i, dv_s, dv_c):
        ws.add_data_validation(dv)

    first, last = 12, 41
    for r in range(first, last + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.border = THIN
            clear_fill(cell)  # cột thành phần / ô dữ liệu: không tô màu
            cell.font = font(False, 14)
            cell.alignment = align("left", "center", True)
        ws.cell(r, 10, prio_formula("H", "I", r))
        ws.cell(r, 13, prio_formula("K", "L", r))
        for prio_col in (10, 13):
            ws.cell(r, prio_col).alignment = align("center", "center", False)
            ws.cell(r, prio_col).font = font(True, 14)
            clear_fill(ws.cell(r, prio_col))
        ws.row_dimensions[r].height = 28

    dv_c.add(f"B{first}:B{last}")
    dv_s.add(f"E{first}:E{last}")
    dv_p.add(f"H{first}:H{last}")
    dv_i.add(f"I{first}:I{last}")
    dv_p.add(f"K{first}:K{last}")
    dv_i.add(f"L{first}:L{last}")

    # Priority: tô màu theo mức A/B/C (nổi bật)
    prio_font = font(True, 14)
    for col in ("J", "M"):
        ws.conditional_formatting.add(
            f"{col}{first}:{col}{last}",
            FormulaRule(formula=[f'${col}{first}="A"'], fill=fill_color(RED), font=prio_font),
        )
        ws.conditional_formatting.add(
            f"{col}{first}:{col}{last}",
            FormulaRule(formula=[f'${col}{first}="B"'], fill=fill_color(ORANGE), font=prio_font),
        )
        ws.conditional_formatting.add(
            f"{col}{first}:{col}{last}",
            FormulaRule(formula=[f'${col}{first}="C"'], fill=fill_color(GREEN), font=prio_font),
        )

    ws.cell(
        last + 2,
        1,
        "* Initial = đánh giá lúc nhận diện. Current = đánh giá lại sau khi thực hiện ứng phó (residual). "
        "Priority tự tính từ P×I theo ma trận phía trên — không sửa tay.",
    )
    ws.cell(last + 2, 1).font = Font(name="Times New Roman", italic=True, size=12, color="666666")
    ws.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=ncols)

    # --- Analysis ---
    ws = wb.create_sheet("Analysis")
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([18, 16, 16, 16, 16, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.merge_cells("A1:F1")
    ws["A1"] = "RISK ANALYSIS WORKSHEET (mẫu — copy 1 sheet / rủi ro khi fill)"
    ws["A1"].font = font(True, 14, WHITE)
    for c in range(1, 7):
        ws.cell(1, c).fill = fill_color(NAVY)

    # simpler layout
    ws["A3"] = "Risk ID"
    ws["B3"] = ""
    ws["C3"] = "Category"
    ws["D3"] = ""
    ws["A4"] = "Probability"
    ws["B4"] = ""
    ws["C4"] = "Impact"
    ws["D4"] = ""
    ws["A5"] = "Priority"
    ws["B5"] = '=IF(OR(B4="",D4=""),"",IF(AND(B4="High",D4="Big"),"A",IF(AND(B4="High",D4="Middle"),"A",IF(AND(B4="High",D4="Small"),"B",IF(AND(B4="Middle",D4="Big"),"A",IF(AND(B4="Middle",D4="Middle"),"B",IF(AND(B4="Middle",D4="Small"),"C",IF(AND(B4="Low",D4="Big"),"B",IF(AND(B4="Low",D4="Middle"),"C",IF(AND(B4="Low",D4="Small"),"C",""))))))))))'
    ws["C5"] = "Strategy"
    ws["D5"] = ""
    ws["A6"] = "Owner"
    ws["B6"] = ""
    for r in range(3, 7):
        for c in range(1, 5):
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(False, 14)
        ws.cell(r, 1).fill = fill_color(GRAY)
        ws.cell(r, 1).font = font(True, 14)
        ws.cell(r, 3).fill = fill_color(GRAY)
        ws.cell(r, 3).font = font(True, 14)
        clear_fill(ws.cell(r, 2))
        clear_fill(ws.cell(r, 4))

    # Priority Analysis: tô màu theo mức
    for rule_val, color in (("A", RED), ("B", ORANGE), ("C", GREEN)):
        ws.conditional_formatting.add(
            "B5",
            FormulaRule(formula=[f'$B$5="{rule_val}"'], fill=fill_color(color), font=font(True, 14)),
        )

    for title, start in [("Driver", 8), ("Risk statement", 11), ("Trigger", 14), ("Response actions", 17), ("Follow-up / Contingency", 20)]:
        ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=6)
        ws.cell(start, 1, title).font = font(True, 14, WHITE)
        for c in range(1, 7):
            ws.cell(start, c).fill = fill_color(NAVY2)
        ws.merge_cells(start_row=start + 1, start_column=1, end_row=start + 2, end_column=6)
        for r in range(start + 1, start + 3):
            for c in range(1, 7):
                ws.cell(r, c).border = THIN
                clear_fill(ws.cell(r, c))
                ws.cell(r, c).font = font(False, 14)

    dv_pa = DataValidation(type="list", formula1="=Values!$A$2:$A$4", allow_blank=True)
    dv_ia = DataValidation(type="list", formula1="=Values!$B$2:$B$4", allow_blank=True)
    dv_sa = DataValidation(type="list", formula1="=Values!$C$2:$C$5", allow_blank=True)
    ws.add_data_validation(dv_pa)
    ws.add_data_validation(dv_ia)
    ws.add_data_validation(dv_sa)
    dv_pa.add("B4")
    dv_ia.add("D4")
    dv_sa.add("D5")

    # order
    order = ["Cover", "Plan", "Values", "Register", "Analysis"]
    for i, name in enumerate(order):
        wb.move_sheet(name, offset=i - wb.sheetnames.index(name))

    wb.save(path)
    print(f"Đã tạo template {path}")
    return path


# ---------------------------------------------------------------------------
# Fill from JSON
# ---------------------------------------------------------------------------
def fill_from_json(data: dict, template: Path, out: Path) -> None:
    shutil.copy2(template, out)
    wb = load_workbook(out)

    project = data.get("project") or {}
    ws = wb["Cover"]
    ws["C6"] = as_plain_text(project.get("name")) or as_plain_text(data.get("project_name"))
    ws["C7"] = as_plain_text(project.get("code"))
    ws["C8"] = as_plain_text(project.get("company"))
    ws["C9"] = as_plain_text(project.get("pm"))
    if as_plain_text(project.get("date")):
        ws["C10"] = as_plain_text(project.get("date"))
    if as_plain_text(project.get("version")):
        ws["C11"] = as_plain_text(project.get("version"))

    plan = data.get("plan") or {}
    ws = wb["Plan"]
    # section value cells at rows 5, 8, 11, 14, 17, 20
    mapping = [
        (5, "methodology"),
        (8, "roles"),
        (11, "budgeting"),
        (14, "timing"),
        (17, "risk_categories"),
    ]
    for row, key in mapping:
        text = as_plain_text(plan.get(key))
        if text:
            ws.cell(row, 1, text)
            ws.cell(row, 1).alignment = align("left", "top", True)

    monitor = data.get("monitor") or {}
    qs = monitor.get("review_questions") or []
    cadence = as_plain_text(monitor.get("cadence"))
    mon_lines = []
    if cadence:
        mon_lines.append(f"Cadence: {cadence}")
    if isinstance(qs, list):
        mon_lines.extend(str(q) for q in qs)
    elif qs:
        mon_lines.append(str(qs))
    if mon_lines:
        ws.cell(20, 1, "\n".join(mon_lines))
        ws.cell(20, 1).alignment = align("left", "top", True)

    # Categories on Values sheet (extend list)
    cats = plan.get("risk_categories")
    cat_list = []
    if isinstance(cats, dict):
        v = cats.get("value")
        if isinstance(v, list):
            cat_list = [str(x) for x in v]
        elif v:
            cat_list = [as_plain_text(v)]
    elif isinstance(cats, list):
        cat_list = [str(x) for x in cats]
    if cat_list:
        ws_v = wb["Values"]
        for i, cat in enumerate(cat_list, 2):
            ws_v.cell(i, 4, cat).border = THIN

    # Register: row 12+ ; Initial P/I/Pri = H/I/J ; Current P/I/Pri = K/L/M
    ws = wb["Register"]
    risks = data.get("risks") or []
    register_data_start = 12
    for i, risk in enumerate(risks[:30]):
        r = register_data_start + i
        rid = as_plain_text(risk.get("id")) or f"R-{i+1:03d}"
        prob = as_plain_text(risk.get("probability"))
        impact = as_plain_text(risk.get("impact"))
        prio = as_plain_text(risk.get("priority")) or compute_priority(prob, impact)
        prob_c = as_plain_text(risk.get("probability_current"))
        impact_c = as_plain_text(risk.get("impact_current"))
        prio_c = as_plain_text(risk.get("priority_current")) or (
            compute_priority(prob_c, impact_c) if prob_c and impact_c else ""
        )

        crs = []
        for k in ("contingency", "residual", "secondary"):
            t = as_plain_text(risk.get(k))
            if t:
                crs.append(f"{k.capitalize()}: {t}")

        # A..G identity/response ; H/I initial P/I ; K/L current P/I ; N/O notes
        # J and M keep Excel formulas
        cells = {
            1: rid,
            2: as_plain_text(risk.get("category")),
            3: as_plain_text(risk.get("driver")),
            4: as_plain_text(risk.get("risk")),
            5: as_plain_text(risk.get("strategy")),
            6: as_plain_text(risk.get("response_actions")),
            7: as_plain_text(risk.get("owner")),
            8: prob,
            9: impact,
            11: prob_c,
            12: impact_c,
            14: " | ".join(crs),
            15: as_plain_text(risk.get("evidence")) or as_plain_text(risk.get("comments")),
        }
        for c, val in cells.items():
            if val:
                cell = ws.cell(r, c, val)
                cell.alignment = align("left", "center", True)
                cell.border = THIN
                cell.font = font(False, 14)
                clear_fill(cell)
        for prio_col in (10, 13):
            ws.cell(r, prio_col).font = font(True, 14)
            clear_fill(ws.cell(r, prio_col))

        expected = compute_priority(prob, impact)
        if prio and expected and prio != expected:
            print(
                f"Cảnh báo: {rid} Initial priority JSON={prio} nhưng ma trận={expected}.",
                file=sys.stderr,
            )
        if prob_c and impact_c and prio_c:
            exp_c = compute_priority(prob_c, impact_c)
            if prio_c != exp_c:
                print(
                    f"Cảnh báo: {rid} Current priority JSON={prio_c} nhưng ma trận={exp_c}.",
                    file=sys.stderr,
                )

    # Analysis: giữ sheet "Analysis" làm mẫu trống; mỗi rủi ro → 1 sheet riêng
    fill_analysis_sheets(wb, risks)

    wb.save(out)
    print(f"Đã ghi {out}")


def safe_sheet_title(rid: str, used: set[str]) -> str:
    """Excel sheet name ≤31 ký tự, không chứa \\ / ? * [ ]."""
    bad = '\\/?*[]'
    title = "".join("_" if ch in bad else ch for ch in rid).strip() or "Risk"
    title = title[:31]
    base = title
    n = 2
    while title in used:
        suffix = f"_{n}"
        title = (base[: 31 - len(suffix)] + suffix)[:31]
        n += 1
    used.add(title)
    return title


def write_analysis_sheet(wa, risk: dict, *, title_suffix: str = "") -> None:
    rid = as_plain_text(risk.get("id"))
    heading = "RISK ANALYSIS WORKSHEET"
    if rid:
        heading = f"{heading} — {rid}"
    if title_suffix:
        heading = f"{heading} {title_suffix}"
    wa["A1"] = heading
    wa["A1"].font = font(True, 14, WHITE)
    wa["B3"] = rid
    wa["D3"] = as_plain_text(risk.get("category"))
    wa["B4"] = as_plain_text(risk.get("probability"))
    wa["D4"] = as_plain_text(risk.get("impact"))
    wa["D5"] = as_plain_text(risk.get("strategy"))
    wa["B6"] = as_plain_text(risk.get("owner"))
    wa["A9"] = as_plain_text(risk.get("driver"))
    wa["A12"] = as_plain_text(risk.get("risk"))
    trigger = as_plain_text(risk.get("trigger"))
    if trigger:
        wa["A15"] = trigger
    wa["A18"] = as_plain_text(risk.get("response_actions"))
    follow = " | ".join(
        filter(
            None,
            [
                as_plain_text(risk.get("contingency")),
                as_plain_text(risk.get("residual")),
                as_plain_text(risk.get("secondary")),
            ],
        )
    )
    if follow:
        wa["A21"] = follow
    for addr in ("B3", "D3", "B4", "D4", "D5", "B6", "A9", "A12", "A15", "A18", "A21"):
        cell = wa[addr]
        if cell.value not in (None, ""):
            cell.font = font(False, 14)
            clear_fill(cell)


def fill_analysis_sheets(wb, risks: list) -> None:
    """Tạo 1 phiếu Analysis / rủi ro; giữ sheet Analysis làm mẫu trống."""
    if "Analysis" not in wb.sheetnames:
        return
    template_ws = wb["Analysis"]
    used_titles = set(wb.sheetnames)
    # Xóa sheet Analysis cũ từng rủi ro (refill)
    stale = [
        name
        for name in list(wb.sheetnames)
        if name != "Analysis"
        and (name.startswith("R-") or name.startswith("Analysis-"))
    ]
    for name in stale:
        del wb[name]
        used_titles.discard(name)

    for i, risk in enumerate(risks[:30]):
        rid = as_plain_text(risk.get("id")) or f"R-{i + 1:03d}"
        title = safe_sheet_title(rid, used_titles)
        wa = wb.copy_worksheet(template_ws)
        wa.title = title
        write_analysis_sheet(wa, {**risk, "id": rid})

    # Mẫu trống: ghi chú số phiếu đã tạo
    n = min(len(risks), 30)
    template_ws["A1"] = (
        f"RISK ANALYSIS WORKSHEET (mẫu trống — đã tạo {n} phiếu: xem các sheet R-xxx)"
        if n
        else "RISK ANALYSIS WORKSHEET (mẫu trống)"
    )
    template_ws["A1"].font = font(True, 14, WHITE)


def main() -> int:
    parser = argparse.ArgumentParser(description="Fill / build Risk Management workbook")
    parser.add_argument("--build-template", action="store_true", help="Tạo lại template xlsx")
    parser.add_argument("--force", action="store_true", help="Fill dù chưa approved")
    args = parser.parse_args()

    if args.build_template:
        build_template(TEMPLATE_PATH)
        return 0

    if not TEMPLATE_PATH.exists():
        build_template(TEMPLATE_PATH)

    if not JSON_PATH.exists():
        print(f"Thiếu {JSON_PATH}", file=sys.stderr)
        return 1

    data = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    if data.get("status") != "approved" and not args.force:
        print(
            f"Từ chối fill: status={data.get('status')!r}. "
            "Chỉ chạy khi status=approved (hoặc --force).",
            file=sys.stderr,
        )
        return 2

    fill_from_json(data, TEMPLATE_PATH, OUT_PATH)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
